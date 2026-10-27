"""
MEGADAPT - Evaluación NLP de la codificación cualitativa, v4

Compara las etiquetas que produce el pipeline (megadapt_qualitative_pipeline_v7.x)
contra la codificación humana oficial: hoja «Integracion» de
2_Análisis_Cualitativo_Integración_formato_estandarizado.xlsx.

Cambios respecto a v3 (megadapt_eval_nlp_v3.py):
  A. Recibe carpetas de corrida (o la carpeta outputs_v<versión>/ completa) en lugar
     de rutas a CSV. De cada corrida lee su experiment_config.json, así que el
     experimento, el modelo, el prompt, la temperatura y el tope quedan como columnas.
  B. Las calificaciones de cada corrida se guardan DENTRO de su carpeta, en eval_v4/,
     con un eval_config.json que dice con qué se evaluó. Si ya existe una evaluación
     con la misma configuración se reutiliza; si la configuración cambió, el script
     se detiene (usar --reevaluar para reemplazarla).
  C. El comparativo (varias corridas juntas) se guarda en
     outputs_v<versión>/comparativos/<fecha_hora>/, con su comparativo_config.json.
  D. Lectura del Excel humano corregida:
       - tabla de equivalencias de claves (I03262015_OTR = I032615_OTR) y sufijo
         GOV = GOB; ya no hay emparejamiento aproximado por prefijo;
       - las filas de bloque (RESUMEN o nombre de codificador) ya no se pegan a la
         entrevista anterior;
       - las celdas se parten solo por salto de línea (ya no por "/" ni ";");
       - los renglones que terminan en ":" son encabezados y no cuentan como
         etiqueta; en Acciones, el encabezado es el actor de las acciones de abajo;
       - se quita la numeración al inicio del renglón ("2. ...").
  E. Campos vacíos: el F1 se calcula solo donde humano y modelo tienen etiquetas.
     Aparte se reporta la presencia/ausencia (ambos, ambos vacíos, solo humano,
     solo modelo).
  F. Estadísticos de distribución (n, media, DE, mín, Q1, mediana, Q3, máx) y número
     promedio de etiquetas humano vs. modelo.
  G. Las repeticiones (E1, E1_rep2, E1_rep3) se agrupan en una condición: por
     entrevista se promedian sus repeticiones (n = entrevistas, no corridas). El
     ruido entre repeticiones se reporta aparte.
  H. Comparación pareada, entrevista por entrevista, contra la referencia (E1) y
     contra la condición anterior con el mismo prompt, con prueba de Wilcoxon.
  I. BERTScore con un solo modelo fijo (bert-base-multilingual-cased), cargado una
     sola vez. Se quita el modo --enrich.
  J. Piso aleatorio: las etiquetas del modelo de cada entrevista se califican también
     contra la codificación humana de todas las demás entrevistas. Se reporta el piso,
     la ganancia sobre el piso (F1 − piso) y el rango (1 = su propia codificación es
     la más parecida). La comparación pareada se hace sobre F1 y sobre la ganancia.
  K. Parecido entre entrevistas: humano contra humano y modelo contra modelo, y su
     diferencia (positiva = el modelo hace que las entrevistas se parezcan más entre sí
     de lo que se parecen según la codificación humana).

La matemática de la calificación es la misma de v3: para cada etiqueta humana se
toma su mejor pareja en el modelo y se promedian (recall); para cada etiqueta del
modelo, su mejor pareja humana (precisión); F1 es la media armónica de las dos.

Uso (ver más ejemplos con --help):
  python megadapt_eval_nlp_v4.py --runs outputs_v7.3

Requiere: pandas, openpyxl, python-dotenv, bert-score (BERTScore), sacrebleu
(BLEU y chrF) y scipy (prueba de Wilcoxon). El Excel humano se toma de --human o de
MEGADAPT_HUMANO en el archivo .env.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import logging
import math
import os
import re
import shutil
import unicodedata
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd
from dotenv import load_dotenv

load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ──────────────────────────────────────────────────────────────────────────────
# 0. Constantes y reglas de evaluación
# ──────────────────────────────────────────────────────────────────────────────
# Versión de las reglas de evaluación: define la subcarpeta eval_v<versión>/ de cada
# corrida y queda anotada en cada eval_config.json. Súbela si cambias las reglas.
EVAL_VERSION = "4"
EVAL_DIRNAME = f"eval_v{EVAL_VERSION}"

BERTSCORE_MODEL_DEFAULT = "bert-base-multilingual-cased"
METRICAS = ("bertscore", "chrf", "bleu")          # orden de prioridad de la principal
METRICA_NOMBRE = {"bertscore": "BERTScore", "chrf": "chrF++", "bleu": "BLEU"}

CAMPOS = ["preocupacion", "causa", "consecuencia", "accion"]
CAMPO_NOMBRE = {
    "preocupacion": "Preocupación",
    "causa": "Causas",
    "consecuencia": "Consecuencias",
    "accion": "Acciones",
    "TOTAL": "TOTAL",
}
LLM_CSV_COLS = {
    "preocupacion": "Preocupacion_principal_acerca_del_agua",
    "causa": "Causas_principales",
    "consecuencia": "Consecuencias",
    "accion": "Acciones_y_Actores_que_realizan_dichas_acciones",
}
HUMAN_HEADER_KEYS = {
    "preocupacion": "preocupaci",
    "causa": "causas principales",
    "consecuencia": "consecuencias",
    "accion": "acciones",
}
NULL_ITEMS = {"ninguno", "ninguna", "none", "na", "n a", "n/a", "no aplica", "nan", "-"}

# Claves que se escriben distinto en el Excel humano y en las transcripciones.
# clave humana -> clave del modelo. Cada entrada debe estar justificada.
EQUIVALENCIAS_CLAVES = {
    "I03262015_OTR": "I032615_OTR",   # el Excel humano escribe el año con 4 dígitos
}
# Sufijos de actor que significan lo mismo (Gobierno).
SUFIJOS_EQUIVALENTES = {"GOV": "GOB"}

# Un renglón que termina en ":" y tiene a lo más estos caracteres es un encabezado.
ENCABEZADO_MAX_CHARS = 40

ORDEN_PROMPTS = ["p1_baseline", "p2_definiciones", "p3_perfil"]
PRESENCIAS = ["ambos", "ambos_vacios", "solo_humano", "solo_modelo"]
PRESENCIA_NOMBRE = {
    "ambos": "Ambos con etiquetas",
    "ambos_vacios": "Ambos vacíos",
    "solo_humano": "Solo humano (modelo vacío)",
    "solo_modelo": "Solo modelo (humano vacío)",
}

REGLAS = [
    "Claves: se comparan sin acentos, mayúsculas/minúsculas, espacios ni guiones bajos; "
    "GOV = GOB; equivalencias explícitas en EQUIVALENCIAS_CLAVES. Sin emparejamiento aproximado.",
    "Excel humano: una fila sin clave y sin Sector continúa la entrevista anterior; una fila "
    "sin clave pero con Sector (RESUMEN o nombre de codificador) es de bloque y se omite.",
    "Excel humano: cada celda se parte solo por salto de línea (no por '/' ni ';').",
    f"Excel humano: un renglón que termina en ':' (hasta {ENCABEZADO_MAX_CHARS} caracteres) es "
    "encabezado y no cuenta como etiqueta; en Acciones se antepone como actor a las acciones de "
    "abajo, dentro de la misma celda.",
    "Excel humano: se quita la numeración al inicio del renglón; se descartan etiquetas de menos "
    "de 3 caracteres, las vacías ('Ninguno', 'N/A'...) y las repetidas dentro de un campo.",
    "Modelo: cada celda del CSV se parte por salto de línea; se descartan repetidas.",
    "F1 solo donde humano y modelo tienen etiquetas; los demás casos se cuentan en la tabla de "
    "presencia.",
    "Calificación de un campo: recall = promedio, sobre las etiquetas humanas, de su mejor "
    "pareja en el modelo; precisión = promedio, sobre las etiquetas del modelo, de su mejor "
    "pareja humana; F1 = media armónica. BLEU y chrF se dividen entre 100.",
    "TOTAL de una entrevista: promedio de los campos que tienen F1.",
    "Piso aleatorio: las etiquetas del modelo de cada entrevista se califican también contra la "
    "codificación humana de cada una de las otras entrevistas; el piso es el promedio de esas "
    "calificaciones cruzadas y la ganancia es F1 propio menos piso. Rango: lugar que ocupa la "
    "propia codificación humana entre todas (1 = es la más parecida).",
    "Parecido entre entrevistas: F1 de la codificación humana de cada entrevista contra la de cada "
    "una de las otras, y lo mismo con las etiquetas del modelo; solo entre entrevistas con F1 en el "
    "campo. Diferencia = modelo − humano.",
]

# ──────────────────────────────────────────────────────────────────────────────
# 1. Utilidades
# ──────────────────────────────────────────────────────────────────────────────
def strip_accents(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", str(s))
        if unicodedata.category(c) != "Mn"
    )


def dedupe(items) -> list[str]:
    seen, out = set(), []
    for x in items:
        k = x.lower()
        if x and k not in seen:
            seen.add(k)
            out.append(x)
    return out


def sha256_file(path: Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 20), b""):
            h.update(chunk)
    return h.hexdigest()


def _norm_key(code: str) -> str:
    c = strip_accents(str(code)).upper().strip()
    c = re.sub(r"[^A-Z0-9]", "", c)
    c = re.sub(r"\d+(RA|DA|TA|ER)?PARTE", "", c)
    for a, b in SUFIJOS_EQUIVALENTES.items():
        if c.endswith(a):
            c = c[: -len(a)] + b
    return c


_EQUIV_NORM = {_norm_key(k): _norm_key(v) for k, v in EQUIVALENCIAS_CLAVES.items()}


def human_key(code: str) -> str:
    k = _norm_key(code)
    return _EQUIV_NORM.get(k, k)


def model_key(code: str) -> str:
    return _norm_key(code)


def describe(values) -> dict:
    """n, media, DE, mín, Q1, mediana, Q3, máx (cuartiles = QUARTILE.INC de Excel)."""
    s = pd.Series(values, dtype="float64").dropna()
    n = len(s)
    if n == 0:
        return {k: (0 if k == "n" else math.nan)
                for k in ("n", "media", "de", "min", "q1", "mediana", "q3", "max")}
    return {
        "n": n,
        "media": round(s.mean(), 4),
        "de": round(s.std(ddof=1), 4) if n > 1 else math.nan,
        "min": round(s.min(), 4),
        "q1": round(s.quantile(0.25), 4),
        "mediana": round(s.median(), 4),
        "q3": round(s.quantile(0.75), 4),
        "max": round(s.max(), 4),
    }


# ──────────────────────────────────────────────────────────────────────────────
# 2. Lectura de la codificación humana
# ──────────────────────────────────────────────────────────────────────────────
_RE_NUM = re.compile(r"^\(?\d{1,2}\s*[.)]\s+")


def parse_human_cell(text, campo: str) -> tuple[list[str], list[str]]:
    """Parte una celda humana en etiquetas. Regresa (etiquetas, encabezados)."""
    if not isinstance(text, str):
        return [], []
    items, heads, actor = [], [], None
    for line in text.split("\n"):
        line = re.sub(r"\s+", " ", line).strip()
        if not line:
            continue
        line = _RE_NUM.sub("", line).strip(" \t•*-")
        if line.endswith(":") and len(line) <= ENCABEZADO_MAX_CHARS:
            head = line.rstrip(":").strip()
            heads.append(head)
            if campo == "accion":
                actor = head
            continue
        line = line.strip(" \t.,;:-•*")
        if len(line) < 3 or strip_accents(line).lower().strip() in NULL_ITEMS:
            continue
        items.append(f"{actor}: {line}" if actor else line)
    return items, heads


@dataclass
class HumanCoding:
    path: Path
    sheet: str
    sha256: str
    data: dict            # clave humana -> {campo: [etiquetas]}
    heads: dict           # clave humana -> {campo: [encabezados]}
    filas_omitidas: list  # (fila de Excel, sector) de las filas de bloque con contenido


def load_human(path: Path, sheet: str) -> HumanCoding:
    raw = pd.read_excel(path, sheet_name=sheet, header=None, dtype=str)
    hdr = None
    for i in range(min(10, len(raw))):
        if any(isinstance(v, str) and strip_accents(v).lower().strip().startswith("codigo")
               for v in raw.iloc[i]):
            hdr = i
            break
    if hdr is None:
        raise SystemExit(f"No encontré la fila de encabezados ('Código...') en la hoja {sheet}.")
    header = [strip_accents(str(v)).lower().strip() if isinstance(v, str) else ""
              for v in raw.iloc[hdr]]
    code_idx = next(i for i, h in enumerate(header) if h.startswith("codigo"))
    sector_idx = next((i for i, h in enumerate(header) if h.startswith("sector")), None)
    if sector_idx is None:
        log.warning("No encontré la columna 'Sector': no se pueden detectar filas de bloque.")
    idx = {}
    for campo, key in HUMAN_HEADER_KEYS.items():
        idx[campo] = next((i for i, h in enumerate(header) if key in h), None)
        if idx[campo] is None:
            raise SystemExit(f"No encontré la columna '{key}' en el Excel humano: {header}")
    extra_acc = list(range(idx["accion"] + 1, raw.shape[1]))
    label_cols = sorted(set(idx.values()) | set(extra_acc))

    data, heads, omitidas = {}, {}, []
    current = None
    for i in range(hdr + 1, len(raw)):
        row = raw.iloc[i]
        code = row.iloc[code_idx]
        sector = row.iloc[sector_idx] if sector_idx is not None else None
        if isinstance(code, str) and code.strip():
            current = code.strip()
            if current in data:
                log.warning(f"Clave humana repetida en el Excel: {current} (se juntan sus filas)")
            data.setdefault(current, {c: [] for c in CAMPOS})
            heads.setdefault(current, {c: [] for c in CAMPOS})
        elif isinstance(sector, str) and sector.strip():
            # Fila de bloque: RESUMEN o nombre de codificador. No es de ninguna entrevista.
            if any(isinstance(row.iloc[j], str) and row.iloc[j].strip() for j in label_cols):
                omitidas.append((i + 1, sector.strip()))
            current = None
            continue
        if current is None:
            continue
        for campo in CAMPOS:
            cells = [row.iloc[idx[campo]]]
            if campo == "accion":
                cells += [row.iloc[j] for j in extra_acc]
            for cell in cells:
                its, hds = parse_human_cell(cell, campo)
                data[current][campo] += its
                heads[current][campo] += hds
    for k in data:
        for campo in CAMPOS:
            data[k][campo] = dedupe(data[k][campo])
    return HumanCoding(Path(path), sheet, sha256_file(Path(path)), data, heads, omitidas)


# ──────────────────────────────────────────────────────────────────────────────
# 3. Corridas del pipeline
# ──────────────────────────────────────────────────────────────────────────────
@dataclass
class Run:
    carpeta: Path
    config: dict
    csv: Path
    experimento: str          # etiqueta completa (ej. E1_rep2)
    base: str                 # etiqueta sin _repN (ej. E1)
    prompt: str
    eval_dir: Path = field(init=False)
    nombre: str = ""          # ruta relativa para mostrar
    eval_cfg: dict = field(default_factory=dict)
    df: Optional[pd.DataFrame] = None
    estado: str = ""          # nueva / reutilizada

    def __post_init__(self):
        self.eval_dir = self.carpeta / EVAL_DIRNAME


def _outputs_root(path: Path) -> Optional[Path]:
    for p in [path, *path.parents]:
        if p.name.startswith("outputs_v"):
            return p
    return None


def find_runs(paths: list[str]) -> list[Run]:
    folders = []
    for p in paths:
        p = Path(p).resolve()
        if not p.exists():
            raise SystemExit(f"--runs: no existe {p}")
        if (p / "experiment_config.json").exists():
            folders.append(p)
        else:
            folders += [c.parent for c in p.rglob("experiment_config.json")]
    runs, seen = [], set()
    for f in sorted(set(folders)):
        if f in seen or "comparativos" in f.parts or any(x.startswith("eval_v") for x in f.parts):
            continue
        seen.add(f)
        cfg = json.loads((f / "experiment_config.json").read_text(encoding="utf-8"))
        prompt = cfg.get("prompt_variant", "")
        csv = f / f"analysis_v7_{prompt}.csv"
        if not csv.exists():
            otros = sorted(f.glob("analysis_*.csv"))
            if not otros:
                log.warning(f"Se omite {f}: no tiene CSV de resultados (¿corrida fallida?).")
                continue
            csv = otros[0]
        # Las corridas anteriores a --esquema usaron el esquema con largo fijo ("4 o 5 palabras")
        cfg.setdefault("esquema", "con_largo")
        etiqueta = cfg.get("experimento") or f.name
        base = re.sub(r"_rep\d+$", "", etiqueta)
        runs.append(Run(carpeta=f, config=cfg, csv=csv, experimento=etiqueta,
                        base=base, prompt=prompt))
    if not runs:
        raise SystemExit("No encontré corridas (carpetas con experiment_config.json y CSV).")
    for r in runs:
        root = _outputs_root(r.carpeta)
        r.nombre = str(r.carpeta.relative_to(root)) if root else r.carpeta.name
    return runs


def load_model_csv(path: Path) -> dict:
    df = pd.read_csv(path, dtype=str, encoding="utf-8-sig").fillna("")
    faltan = [c for c in ["Id_entrevista", *LLM_CSV_COLS.values()] if c not in df.columns]
    if faltan:
        raise SystemExit(f"{path}: faltan columnas {faltan}")
    out = {}
    for _, r in df.iterrows():
        d = out.setdefault(str(r["Id_entrevista"]).strip(), {c: [] for c in CAMPOS})
        for campo in CAMPOS:
            d[campo] += [x.strip() for x in str(r[LLM_CSV_COLS[campo]]).split("\n") if x.strip()]
    for d in out.values():
        for campo in CAMPOS:
            d[campo] = dedupe(d[campo])
    return out


def match_ids(human: dict, model: dict) -> tuple[dict, list, list]:
    """Empareja por clave normalizada. Regresa ({clave humana: [ids modelo]}, sin par modelo, sin par humano)."""
    hk = {}
    for h in human:
        k = human_key(h)
        if k in hk:
            log.warning(f"Dos claves humanas equivalen a {k}: {hk[k]} y {h}")
        hk[k] = h
    mk = defaultdict(list)
    for m in model:
        mk[model_key(m)].append(m)
    pairs = {hk[k]: mids for k, mids in mk.items() if k in hk}
    sin_par_modelo = sorted(m for k, mids in mk.items() if k not in hk for m in mids)
    sin_par_humano = sorted(h for k, h in hk.items() if k not in mk)
    for h, mids in pairs.items():
        if len(mids) > 1:
            log.warning(f"Varias filas del modelo para {h}: {mids} (se juntan sus etiquetas)")
    return pairs, sin_par_modelo, sin_par_humano


# ──────────────────────────────────────────────────────────────────────────────
# 4. Métricas (misma matemática que v3)
# ──────────────────────────────────────────────────────────────────────────────
class SacreSim:
    """BLEU o chrF++ de sacrebleu, en escala 0-100 (se divide entre 100 al final, como v3)."""
    scale = 100.0

    def __init__(self, kind: str):
        from sacrebleu.metrics import BLEU, CHRF
        self.m = BLEU(smooth_method="exp", effective_order=True) if kind == "bleu" else CHRF(word_order=2)
        self.cache = {}

    def prefetch(self, pairs):
        pass

    def sim(self, cand: str, ref: str) -> float:
        key = (cand, ref)
        if key not in self.cache:
            self.cache[key] = self.m.sentence_score(cand, [ref]).score
        return self.cache[key]


class BertScoreSim:
    """BERTScore F de cada par (candidato = modelo, referencia = humano). Modelo cargado una vez."""
    scale = 1.0

    def __init__(self, model_type: str, device: str):
        from bert_score import BERTScorer
        log.info(f"Cargando BERTScore ({model_type}, {device})...")
        self.scorer = BERTScorer(model_type=model_type, lang="es", device=device)
        self.cache = {}

    def prefetch(self, pairs):
        todo = sorted({p for p in pairs if p not in self.cache})
        if not todo:
            return
        cands, refs = zip(*todo)
        _, _, F = self.scorer.score(list(cands), list(refs), batch_size=64)
        for p, f in zip(todo, F.tolist()):
            self.cache[p] = f

    def sim(self, cand: str, ref: str) -> float:
        if (cand, ref) not in self.cache:
            self.prefetch([(cand, ref)])
        return self.cache[(cand, ref)]


def make_sim(metric: str, bertscore_model: str, device: str):
    if metric == "bertscore":
        return BertScoreSim(bertscore_model, device)
    return SacreSim(metric)


def set_scores(h_items: list[str], m_items: list[str], sim) -> dict:
    """Recall, precisión y F1 de un campo (misma matemática que v3)."""
    matrix = [[sim.sim(m, h) for m in m_items] for h in h_items]
    n_h, n_m = len(h_items), len(m_items)
    recall = sum(max(row) for row in matrix) / n_h / sim.scale
    precision = sum(max(matrix[i][j] for i in range(n_h)) for j in range(n_m)) / n_m / sim.scale
    f1 = 2 * recall * precision / (recall + precision) if (recall + precision) > 0 else 0.0
    return {"recall": round(recall, 4), "precision": round(precision, 4), "f1": round(f1, 4)}


def presencia(h_items, m_items) -> str:
    if h_items and m_items:
        return "ambos"
    if not h_items and not m_items:
        return "ambos_vacios"
    return "solo_humano" if h_items else "solo_modelo"


# ──────────────────────────────────────────────────────────────────────────────
# 5. Evaluación de una corrida (se guarda en <corrida>/eval_v4/)
# ──────────────────────────────────────────────────────────────────────────────
def evaluate_run(run: Run, human: HumanCoding, metrics: list[str], sims: dict) -> tuple[pd.DataFrame, dict]:
    """Califica cada entrevista contra su codificación humana y, para el piso aleatorio,
    contra la codificación humana de todas las demás entrevistas."""
    model = load_model_csv(run.csv)
    pairs, sin_par_modelo, sin_par_humano = match_ids(human.data, model)
    if sin_par_modelo:
        log.warning(f"[{run.nombre}] entrevistas del modelo SIN par humano (no se califican): {sin_par_modelo}")
    ids = sorted(pairs)
    H = {i: {c: human.data[i][c] for c in CAMPOS} for i in ids}
    M = {i: {c: dedupe(x for mid in pairs[i] for x in model[mid][c]) for c in CAMPOS} for i in ids}

    # Todas las combinaciones modelo(i) × humano(j), incluida la propia (i = j), más
    # humano(i) × humano(j) y modelo(i) × modelo(j) para las medidas de parecido
    for met in metrics:
        sims[met].prefetch(
            [(m, h) for c in CAMPOS for i in ids for m in M[i][c] for j in ids for h in H[j][c]]
            + [(a, b) for c in CAMPOS for X in (H, M) for i in ids for a in X[i][c]
               for j in ids if j != i for b in X[j][c]])

    rows = []
    for met in metrics:
        # F[c][(i, j)] = F1 de las etiquetas del modelo en i contra las humanas de j
        F = {c: {(i, j): set_scores(H[j][c], M[i][c], sims[met])
                 for i in ids if M[i][c] for j in ids if H[j][c]} for c in CAMPOS}
        # Parecido entre entrevistas distintas, solo entre las que tienen F1 en ese campo
        # (mismas parejas para el humano y para el modelo): P[c][i] = (humano, modelo)
        P = {}
        for c in CAMPOS:
            S = [i for i in ids if (i, i) in F[c]]
            P[c] = {}
            for i in S:
                otros = [j for j in S if j != i]
                if otros:
                    ph = sum(set_scores(H[j][c], H[i][c], sims[met])["f1"] for j in otros) / len(otros)
                    pm = sum(set_scores(M[j][c], M[i][c], sims[met])["f1"] for j in otros) / len(otros)
                    P[c][i] = (ph, pm)
        for i in ids:
            # Rango del TOTAL: ¿en cuántas otras entrevistas el TOTAL cruzado supera al propio?
            mejores, cands = 0, 0
            for j in ids:
                if j == i:
                    continue
                comunes = [c for c in CAMPOS if (i, i) in F[c] and (i, j) in F[c]]
                if comunes:
                    cands += 1
                    propio = sum(F[c][(i, i)]["f1"] for c in comunes) / len(comunes)
                    cruzado = sum(F[c][(i, j)]["f1"] for c in comunes) / len(comunes)
                    mejores += cruzado > propio
            for c in CAMPOS:
                h, m = H[i][c], M[i][c]
                row = {
                    "experimento": run.experimento, "corrida": run.nombre,
                    "Id_humano": i, "Id_modelo": "+".join(pairs[i]), "campo": c,
                    "presencia": presencia(h, m), "n_items_humano": len(h), "n_items_modelo": len(m),
                }
                if (i, i) in F[c]:
                    propio = F[c][(i, i)]
                    cruz = [F[c][(i, j)]["f1"] for j in ids if j != i and (i, j) in F[c]]
                    piso = sum(cruz) / len(cruz) if cruz else math.nan
                    ph, pm = P[c].get(i, (math.nan, math.nan))
                    vals = {**propio,
                            "piso": round(piso, 4),
                            "ganancia": round(propio["f1"] - piso, 4) if cruz else math.nan,
                            "rango": 1 + sum(x > propio["f1"] for x in cruz) if cruz else math.nan,
                            "rango_total": 1 + mejores if cands else math.nan,
                            "parecido_humano": round(ph, 4),
                            "parecido_modelo": round(pm, 4),
                            "dif_parecido": round(pm - ph, 4)}
                    row["n_candidatos"] = 1 + len(cruz)
                    row["n_candidatos_total"] = 1 + cands
                else:
                    vals = {k: math.nan for k in ("recall", "precision", "f1", "piso", "ganancia",
                                                  "rango", "rango_total", "parecido_humano",
                                                  "parecido_modelo", "dif_parecido")}
                for k, v in vals.items():
                    row[f"{met}_{k}"] = v
                if met == metrics[0]:
                    rows.append(row)
                else:
                    rows[ids.index(i) * len(CAMPOS) + CAMPOS.index(c)].update(
                        {f"{met}_{k}": v for k, v in vals.items()})
    info = {
        "entrevistas_en_csv": len(model),
        "emparejadas": len(pairs),
        "sin_par_modelo": sin_par_modelo,
        "humanas_sin_par": sin_par_humano,
    }
    return pd.DataFrame(rows), info


def per_interview_total(df: pd.DataFrame, metric: str, key: list[str]) -> pd.DataFrame:
    """TOTAL de una entrevista: promedio de los campos con F1 (y de su piso y ganancia).
    El rango del TOTAL se calculó aparte en evaluate_run (rango_total)."""
    cols = [f"{metric}_{k}" for k in ("recall", "precision", "f1", "piso", "ganancia",
                                      "parecido_humano", "parecido_modelo", "dif_parecido")]
    sub = df.dropna(subset=[f"{metric}_f1"])
    t = sub.groupby(key, as_index=False)[cols].mean()
    extra = sub.groupby(key, as_index=False).agg(
        **{f"{metric}_rango": (f"{metric}_rango_total", "first"),
           "n_candidatos": ("n_candidatos_total", "first")})
    t = t.merge(extra, on=key)
    t["campo"] = "TOTAL"
    return t


def _piso_stats(sub: pd.DataFrame, metric: str) -> dict:
    g = sub[f"{metric}_ganancia"].dropna()
    rg = sub[f"{metric}_rango"].dropna()
    return {
        "piso_media": round(sub[f"{metric}_piso"].mean(), 4),
        "ganancia_media": round(g.mean(), 4),
        "pct_ganancia_positiva": round(100 * (g > 0).mean(), 1) if len(g) else math.nan,
        "rango_medio": round(rg.mean(), 2) if len(rg) else math.nan,
        "pct_rango_1": round(100 * (rg == 1).mean(), 1) if len(rg) else math.nan,
        "candidatos": round(sub["n_candidatos"].mean(), 1) if "n_candidatos" in sub else math.nan,
        "parecido_humano": round(sub[f"{metric}_parecido_humano"].mean(), 4),
        "parecido_modelo": round(sub[f"{metric}_parecido_modelo"].mean(), 4),
        "dif_parecido": round(sub[f"{metric}_dif_parecido"].mean(), 4),
    }


def summarize_run(df: pd.DataFrame, metrics: list[str]) -> pd.DataFrame:
    rows = []
    for met in metrics:
        tot = per_interview_total(df, met, ["Id_humano"])
        for campo in CAMPOS + ["TOTAL"]:
            sub = tot if campo == "TOTAL" else df[df["campo"] == campo]
            allc = df if campo == "TOTAL" else df[df["campo"] == campo]
            r = {"metrica": met, "campo": campo, **describe(sub[f"{met}_f1"])}
            r["recall_media"] = round(sub[f"{met}_recall"].mean(), 4)
            r["precision_media"] = round(sub[f"{met}_precision"].mean(), 4)
            r.update(_piso_stats(sub, met))
            r["etiquetas_humano_media"] = round(allc["n_items_humano"].mean(), 2)
            r["etiquetas_modelo_media"] = round(allc["n_items_modelo"].mean(), 2)
            for p in PRESENCIAS:
                r[f"n_{p}"] = int((allc["presencia"] == p).sum())
            rows.append(r)
    return pd.DataFrame(rows)


def _firma(run: Run, human: HumanCoding, metrics: list[str], bertscore_model: str) -> dict:
    return {
        "eval_version": EVAL_VERSION,
        "excel_humano_sha256": human.sha256,
        "hoja": human.sheet,
        "metricas": sorted(metrics),
        "bertscore_model": bertscore_model if "bertscore" in metrics else None,
        "csv_sha256": sha256_file(run.csv),
        "equivalencias_claves": EQUIVALENCIAS_CLAVES,
        "piso": "cruzado: modelo(i) contra humano(j) para toda j distinta de i",
        "parecido": "humano(i) vs humano(j) y modelo(i) vs modelo(j), entrevistas con F1 en el campo",
    }


def evaluate_or_load(run: Run, human: HumanCoding, metrics: list[str], bertscore_model: str,
                     get_sims, reevaluar: bool) -> None:
    firma = _firma(run, human, metrics, bertscore_model)
    cfg_path = run.eval_dir / "eval_config.json"
    if cfg_path.exists():
        prev = json.loads(cfg_path.read_text(encoding="utf-8"))
        if prev.get("firma") == firma and not reevaluar:
            run.df = pd.read_csv(run.eval_dir / "eval_por_entrevista.csv", encoding="utf-8-sig")
            run.eval_cfg = prev
            run.estado = "reutilizada"
            log.info(f"[{run.nombre}] evaluación existente reutilizada.")
            return
        if not reevaluar:
            cambios = [k for k in firma if prev.get("firma", {}).get(k) != firma[k]]
            raise SystemExit(
                f"[{run.nombre}] ya tiene una evaluación {EVAL_DIRNAME} hecha con otra "
                f"configuración (cambió: {', '.join(cambios)}).\n"
                f"Usa --reevaluar para reemplazarla, o sube EVAL_VERSION si cambiaste las reglas."
            )
    log.info(f"[{run.nombre}] evaluando {run.csv.name}...")
    df, info = evaluate_run(run, human, metrics, get_sims())
    run.eval_dir.mkdir(parents=True, exist_ok=True)
    df.to_csv(run.eval_dir / "eval_por_entrevista.csv", index=False, encoding="utf-8-sig")
    summarize_run(df, metrics).to_csv(run.eval_dir / "eval_resumen.csv", index=False, encoding="utf-8-sig")
    cfg = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "script": Path(__file__).name,
        "firma": firma,
        "excel_humano": str(human.path.resolve()),
        "excel_humano_modificado": datetime.fromtimestamp(human.path.stat().st_mtime).isoformat(timespec="seconds"),
        "csv_evaluado": str(run.csv.resolve()),
        "experimento": run.experimento,
        "reglas": REGLAS,
        **info,
    }
    cfg_path.write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")
    run.df, run.eval_cfg, run.estado = df, cfg, "nueva"
    if info["sin_par_modelo"]:
        log.warning(f"[{run.nombre}] {len(info['sin_par_modelo'])} entrevista(s) sin par: revisa EQUIVALENCIAS_CLAVES.")
    log.info(f"[{run.nombre}] {info['emparejadas']} de {info['entrevistas_en_csv']} entrevistas emparejadas.")


# ──────────────────────────────────────────────────────────────────────────────
# 6. Comparativo: condiciones, repeticiones, ruido y comparación pareada
# ──────────────────────────────────────────────────────────────────────────────
VARS_CONDICION = ("script_version", "modelo", "temperatura", "prompt_variant",
                  "max_items_postproceso", "prompt_text", "esquema")


def _exp_num(base: str) -> float:
    m = re.match(r"^[Ee](\d+)", base)
    return float(m.group(1)) if m else math.inf


def _prompt_ord(p: str) -> int:
    return ORDEN_PROMPTS.index(p) if p in ORDEN_PROMPTS else len(ORDEN_PROMPTS)


@dataclass
class Condicion:
    base: str
    prompt: str
    runs: list

    @property
    def key(self):
        return (self.base, self.prompt)

    @property
    def nombre(self):
        return f"{self.base} · {self.prompt}"

    def var(self, k):
        return self.runs[0].config.get(k)

    def orden(self):
        fecha = min(r.config.get("fecha", "") for r in self.runs)
        return (_exp_num(self.base), fecha if _exp_num(self.base) == math.inf else "", _prompt_ord(self.prompt))


def build_conditions(runs: list[Run]) -> list[Condicion]:
    grupos = defaultdict(list)
    for r in runs:
        grupos[(r.base, r.prompt)].append(r)
    conds = []
    for (base, prompt), rs in grupos.items():
        for k in VARS_CONDICION:
            vals = {json.dumps(r.config.get(k), ensure_ascii=False, sort_keys=True) for r in rs}
            if len(vals) > 1:
                raise SystemExit(
                    f"Las corridas de {base} · {prompt} no son repeticiones de lo mismo: "
                    f"'{k}' cambia entre {[r.nombre for r in rs]}.\n"
                    f"Usa otra etiqueta de experimento para condiciones distintas."
                )
        rs.sort(key=lambda r: (r.config.get("fecha", ""), r.experimento))
        etiquetas = [r.experimento for r in rs]
        if len(set(etiquetas)) < len(etiquetas):
            log.warning(f"{base} · {prompt}: hay corridas con la misma etiqueta {etiquetas}; "
                        f"conviene usar _rep2, _rep3... para distinguirlas.")
        conds.append(Condicion(base, prompt, rs))
    conds.sort(key=lambda c: c.orden())
    return conds


def check_same_evaluation(runs: list[Run]) -> dict:
    claves = ("eval_version", "excel_humano_sha256", "hoja", "metricas", "bertscore_model", "equivalencias_claves")
    firmas = {json.dumps({k: r.eval_cfg["firma"][k] for k in claves}, sort_keys=True) for r in runs}
    if len(firmas) > 1:
        raise SystemExit("Las corridas se evaluaron con configuraciones distintas (Excel humano, "
                         "métricas o modelo de BERTScore). Reevalúa con --reevaluar para igualarlas.")
    return json.loads(firmas.pop())


def comparisons_for(conds: list[Condicion], referencia: str) -> list[tuple]:
    """(condición, contra, tipo). Anterior = condición previa más reciente con el mismo prompt;
    si no hay, la del experimento inmediato anterior."""
    out = []
    ref_conds = [c for c in conds if c.base == referencia]
    if not ref_conds:
        log.warning(f"No hay corridas del experimento de referencia '{referencia}'.")
    for i, c in enumerate(conds):
        if c.base == referencia:
            continue
        ref = next((r for r in ref_conds if r.prompt == c.prompt), ref_conds[0] if len(ref_conds) == 1 else None)
        previas = [p for p in conds[:i] if p.orden()[0] < c.orden()[0] or
                   (p.orden()[0] == c.orden()[0] == math.inf and p.base != c.base)]
        ant = next((p for p in reversed(previas) if p.prompt == c.prompt), None)
        if ant is None and previas:
            ult = previas[-1].orden()[0]
            cand = [p for p in previas if p.orden()[0] == ult]
            ant = cand[0]
        if ref is not None and ant is not None and ref.key == ant.key:
            out.append((c, ref, "referencia y anterior"))
        else:
            if ref is not None:
                out.append((c, ref, "referencia"))
            if ant is not None:
                out.append((c, ant, "anterior"))
    return out


def _wilcoxon(d: pd.Series) -> float:
    try:
        from scipy.stats import wilcoxon
    except ImportError:
        if not getattr(_wilcoxon, "_avisado", False):
            log.warning("scipy no está instalado: no se calcula la prueba de Wilcoxon (conda install scipy).")
            _wilcoxon._avisado = True
        return math.nan
    d = d[d.abs() > 1e-12]
    if len(d) < 2:
        return math.nan
    try:
        return round(float(wilcoxon(d).pvalue), 4)
    except ValueError:
        return math.nan


def build_comparativo(runs: list[Run], metrics: list[str], referencia: str) -> dict:
    conds = build_conditions(runs)
    cond_of = {id(r): c for c in conds for r in c.runs}
    rows = []
    for r in runs:
        d = r.df.copy()
        d["condicion"] = cond_of[id(r)].nombre
        rows.append(d)
    allr = pd.concat(rows, ignore_index=True)
    principal = next(m for m in METRICAS if m in metrics)
    tablas = {}

    # 6.1 Puntaje por condición × entrevista × campo (promedio de repeticiones)
    pe = []
    for met in metrics:
        f1 = f"{met}_f1"
        valid = allr.dropna(subset=[f1]).copy()
        valid["_r1"] = (valid[f"{met}_rango"] == 1).astype(float)
        g = valid.groupby(["condicion", "Id_humano", "campo"]).agg(
            recall=(f"{met}_recall", "mean"), precision=(f"{met}_precision", "mean"),
            f1=(f1, "mean"), piso=(f"{met}_piso", "mean"), ganancia=(f"{met}_ganancia", "mean"),
            rango=(f"{met}_rango", "mean"), rango1=("_r1", "mean"),
            candidatos=("n_candidatos", "mean"),
            parecido_humano=(f"{met}_parecido_humano", "mean"),
            parecido_modelo=(f"{met}_parecido_modelo", "mean"),
            dif_parecido=(f"{met}_dif_parecido", "mean"),
            repeticiones=(f1, "size")).reset_index()
        tot = g.groupby(["condicion", "Id_humano"]).agg(
            recall=("recall", "mean"), precision=("precision", "mean"), f1=("f1", "mean"),
            piso=("piso", "mean"), ganancia=("ganancia", "mean"),
            parecido_humano=("parecido_humano", "mean"), parecido_modelo=("parecido_modelo", "mean"),
            dif_parecido=("dif_parecido", "mean"),
            repeticiones=("repeticiones", "min")).reset_index()
        # Rango del TOTAL: viene de cada corrida (rango_total) y se promedia entre repeticiones
        rt = valid.groupby(["condicion", "corrida", "Id_humano"]).agg(
            rt=(f"{met}_rango_total", "first"), ct=("n_candidatos_total", "first")).reset_index()
        rt["_r1"] = (rt["rt"] == 1).astype(float)
        rt = rt.groupby(["condicion", "Id_humano"]).agg(
            rango=("rt", "mean"), rango1=("_r1", "mean"), candidatos=("ct", "mean")).reset_index()
        tot = tot.merge(rt, on=["condicion", "Id_humano"], how="left")
        tot["campo"] = "TOTAL"
        gg = pd.concat([g, tot], ignore_index=True)
        gg["metrica"] = met
        pe.append(gg)
    por_ent = pd.concat(pe, ignore_index=True)
    for c in ("recall", "precision", "f1", "piso", "ganancia", "parecido_humano",
              "parecido_modelo", "dif_parecido"):
        por_ent[c] = por_ent[c].round(4)
    por_ent["rango"] = por_ent["rango"].round(2)

    # 6.2 Condiciones (una fila) y distribución (condición × campo)
    cond_rows, dist_rows = [], []
    for c in conds:
        base_row = {
            "condicion": c.nombre, "experimento": c.base, "prompt": c.prompt,
            "modelo": c.var("modelo"), "temperatura": c.var("temperatura"),
            "tope": c.var("max_items_postproceso"), "esquema": c.var("esquema"),
            "repeticiones": len(c.runs),
            "corridas": ", ".join(r.experimento for r in c.runs),
        }
        sub_all = allr[allr["condicion"] == c.nombre]
        for met in metrics:
            crow = {**base_row, "metrica": met}
            for campo in CAMPOS + ["TOTAL"]:
                s = por_ent[(por_ent["condicion"] == c.nombre) & (por_ent["metrica"] == met) &
                            (por_ent["campo"] == campo)]
                desc = describe(s["f1"])
                crow[campo] = desc["media"]
                if campo == "TOTAL":
                    crow["piso_TOTAL"] = round(s["piso"].mean(), 4)
                    crow["ganancia_TOTAL"] = round(s["ganancia"].mean(), 4)
                allc = sub_all if campo == "TOTAL" else sub_all[sub_all["campo"] == campo]
                g = s["ganancia"].dropna()
                dist_rows.append({
                    "condicion": c.nombre, "experimento": c.base, "prompt": c.prompt,
                    "metrica": met, "campo": campo, **desc,
                    "recall_media": round(s["recall"].mean(), 4),
                    "precision_media": round(s["precision"].mean(), 4),
                    "etiquetas_humano_media": round(allc["n_items_humano"].mean(), 2),
                    "etiquetas_modelo_media": round(allc["n_items_modelo"].mean(), 2),
                    "piso_media": round(s["piso"].mean(), 4),
                    "ganancia_media": round(g.mean(), 4),
                    "pct_ganancia_positiva": round(100 * (g > 0).mean(), 1) if len(g) else math.nan,
                    "rango_medio": round(s["rango"].mean(), 2),
                    "pct_rango_1": round(100 * s["rango1"].mean(), 1),
                    "candidatos": round(s["candidatos"].mean(), 1),
                    "parecido_humano": round(s["parecido_humano"].mean(), 4),
                    "parecido_modelo": round(s["parecido_modelo"].mean(), 4),
                    "dif_parecido": round(s["dif_parecido"].mean(), 4),
                })
            cond_rows.append(crow)
    tablas["condiciones"] = pd.DataFrame(cond_rows)
    tablas["distribucion"] = pd.DataFrame(dist_rows)

    # 6.3 Ruido entre repeticiones
    ruido_rows = []
    for c in conds:
        for met in metrics:
            f1 = f"{met}_f1"
            sub = allr[(allr["condicion"] == c.nombre)].dropna(subset=[f1])
            tot = sub.groupby(["corrida", "Id_humano"], as_index=False)[f1].mean()
            tot["campo"] = "TOTAL"
            sub2 = pd.concat([sub[["corrida", "Id_humano", "campo", f1]], tot], ignore_index=True)
            for campo in CAMPOS + ["TOTAL"]:
                s = sub2[sub2["campo"] == campo]
                medias = [round(s[s["corrida"] == r.nombre][f1].mean(), 4) for r in c.runs]
                por_e = s.groupby("Id_humano")[f1].std(ddof=1).dropna()
                ruido_rows.append({
                    "condicion": c.nombre, "metrica": met, "campo": campo,
                    "repeticiones": len(c.runs),
                    "media_por_repeticion": " | ".join(f"{r.experimento}: {m:.4f}" for r, m in zip(c.runs, medias)),
                    "rango_entre_repeticiones": round(max(medias) - min(medias), 4) if len(medias) > 1 else math.nan,
                    "de_entre_repeticiones": round(pd.Series(medias).std(ddof=1), 4) if len(medias) > 1 else math.nan,
                    "variacion_media_por_entrevista": round(por_e.mean(), 4) if len(por_e) else math.nan,
                })
    tablas["ruido"] = pd.DataFrame(ruido_rows)
    rango = {(r["condicion"], r["metrica"], r["campo"]): r["rango_entre_repeticiones"] for r in ruido_rows}

    # 6.4 Presencia / ausencia de etiquetas
    pres_rows = []
    for c in conds:
        sub = allr[allr["condicion"] == c.nombre]
        for campo in CAMPOS + ["TOTAL"]:
            s = sub if campo == "TOTAL" else sub[sub["campo"] == campo]
            r = {"condicion": c.nombre, "campo": campo, "repeticiones": len(c.runs), "casos": len(s)}
            for p in PRESENCIAS:
                r[p] = int((s["presencia"] == p).sum())
            pres_rows.append(r)
    tablas["presencia"] = pd.DataFrame(pres_rows)

    # 6.5 Comparación pareada entrevista por entrevista
    comps = comparisons_for(conds, referencia)
    par_rows, det_rows = [], []
    for c, b, tipo in comps:
        for met, medida, campo in [(m, md, cp) for m in metrics for md in ("f1", "ganancia")
                                   for cp in CAMPOS + ["TOTAL"]]:
                sel = (por_ent["metrica"] == met) & (por_ent["campo"] == campo)
                x = por_ent[sel & (por_ent["condicion"] == c.nombre)].set_index("Id_humano")[medida]
                y = por_ent[sel & (por_ent["condicion"] == b.nombre)].set_index("Id_humano")[medida]
                ids = sorted(set(x.dropna().index) & set(y.dropna().index))
                d = (x.loc[ids] - y.loc[ids]).round(4) if ids else pd.Series(dtype="float64")
                par_rows.append({
                    "condicion": c.nombre, "contra": b.nombre, "tipo": tipo, "metrica": met,
                    "medida": medida, "campo": campo,
                    "n": len(ids),
                    "media_condicion": round(x.loc[ids].mean(), 4) if ids else math.nan,
                    "media_contra": round(y.loc[ids].mean(), 4) if ids else math.nan,
                    "diferencia_media": round(d.mean(), 4) if ids else math.nan,
                    "diferencia_mediana": round(d.median(), 4) if ids else math.nan,
                    "mejoran": int((d > 1e-12).sum()), "empeoran": int((d < -1e-12).sum()),
                    "iguales": int((d.abs() <= 1e-12).sum()),
                    "p_wilcoxon": _wilcoxon(d) if ids else math.nan,
                    "ruido_rango_contra": rango.get((b.nombre, met, campo), math.nan) if medida == "f1" else math.nan,
                    "ruido_rango_condicion": rango.get((c.nombre, met, campo), math.nan) if medida == "f1" else math.nan,
                })
                if met == principal and medida == "f1":
                    for i in ids:
                        det_rows.append({"condicion": c.nombre, "contra": b.nombre, "tipo": tipo,
                                         "Id_humano": i, "campo": campo, "f1_condicion": x.loc[i],
                                         "f1_contra": y.loc[i], "diferencia": round(x.loc[i] - y.loc[i], 4)})
    tablas["pareado"] = pd.DataFrame(par_rows, columns=[
        "condicion", "contra", "tipo", "metrica", "medida", "campo", "n", "media_condicion", "media_contra",
        "diferencia_media", "diferencia_mediana", "mejoran", "empeoran", "iguales", "p_wilcoxon",
        "ruido_rango_contra", "ruido_rango_condicion"])
    tablas["pareado_detalle"] = pd.DataFrame(det_rows, columns=[
        "condicion", "contra", "tipo", "Id_humano", "campo", "f1_condicion", "f1_contra", "diferencia"])

    # 6.6 Matriz por entrevista: filas métrica×entrevista×campo, columnas condición (F1)
    mat = por_ent.pivot_table(index=["metrica", "Id_humano", "campo"], columns="condicion",
                              values="f1").reset_index()
    mat["_ord"] = mat["campo"].map({c: i for i, c in enumerate(CAMPOS + ["TOTAL"])})
    mat["_met"] = mat["metrica"].map({m: i for i, m in enumerate(metrics)})
    mat = mat.sort_values(["_met", "Id_humano", "_ord"]).drop(columns=["_ord", "_met"])
    mat = mat[["metrica", "Id_humano", "campo"] + [c.nombre for c in conds if c.nombre in mat.columns]]
    tablas["por_entrevista"] = mat

    # 6.7 Corridas incluidas
    tablas["corridas"] = pd.DataFrame([{
        "condicion": cond_of[id(r)].nombre, "experimento": r.experimento, "prompt": r.prompt,
        "modelo": r.config.get("modelo"), "temperatura": r.config.get("temperatura"),
        "tope": r.config.get("max_items_postproceso"), "esquema": r.config.get("esquema"),
        "fecha_corrida": r.config.get("fecha"),
        "script_version": r.config.get("script_version"), "carpeta": r.nombre, "csv": r.csv.name,
        "entrevistas_en_csv": r.eval_cfg.get("entrevistas_en_csv"),
        "emparejadas": r.eval_cfg.get("emparejadas"),
        "sin_par_modelo": ", ".join(r.eval_cfg.get("sin_par_modelo", [])),
        "evaluacion": r.estado,
    } for c in conds for r in c.runs])
    return {"conds": conds, "tablas": tablas, "principal": principal, "comparaciones": comps}


# ──────────────────────────────────────────────────────────────────────────────
# 7. Excel comparativo
# ──────────────────────────────────────────────────────────────────────────────
_HEATMAP = ("F8696B", "FFEB84", "63BE7B")
_HEATMAP_THRESH = (0.40, 0.625, 0.85)

ENCABEZADOS = {
    "condicion": "Condición", "experimento": "Experimento", "prompt": "Prompt", "modelo": "Modelo",
    "temperatura": "Temperatura", "tope": "Tope (max-items)", "esquema": "Esquema",
    "repeticiones": "Repeticiones",
    "corridas": "Corridas", "metrica": "Métrica", "campo": "Campo", "n": "n", "media": "Media",
    "de": "DE", "min": "Mín", "q1": "Q1", "mediana": "Mediana", "q3": "Q3", "max": "Máx",
    "recall_media": "Recall medio", "precision_media": "Precisión media",
    "etiquetas_humano_media": "Etiquetas humano (prom.)", "etiquetas_modelo_media": "Etiquetas modelo (prom.)",
    "media_por_repeticion": "Media de cada repetición", "rango_entre_repeticiones": "Rango entre repeticiones",
    "de_entre_repeticiones": "DE entre repeticiones",
    "variacion_media_por_entrevista": "Variación media por entrevista (DE)",
    "casos": "Casos", "ambos": PRESENCIA_NOMBRE["ambos"], "ambos_vacios": PRESENCIA_NOMBRE["ambos_vacios"],
    "solo_humano": PRESENCIA_NOMBRE["solo_humano"], "solo_modelo": PRESENCIA_NOMBRE["solo_modelo"],
    "contra": "Contra", "tipo": "Tipo", "media_condicion": "Media condición", "media_contra": "Media contra",
    "diferencia_media": "Diferencia media", "diferencia_mediana": "Diferencia mediana",
    "mejoran": "Entrevistas que mejoran", "empeoran": "Entrevistas que empeoran", "iguales": "Iguales",
    "p_wilcoxon": "p (Wilcoxon)", "ruido_rango_contra": "Ruido (rango) contra",
    "ruido_rango_condicion": "Ruido (rango) condición", "Id_humano": "Entrevista",
    "f1_condicion": "F1 condición", "f1_contra": "F1 contra", "diferencia": "Diferencia",
    "fecha_corrida": "Fecha de la corrida", "script_version": "Versión pipeline", "carpeta": "Carpeta",
    "csv": "CSV", "entrevistas_en_csv": "Entrevistas en CSV", "emparejadas": "Emparejadas",
    "sin_par_modelo": "Sin par humano", "evaluacion": "Evaluación",
    "piso_TOTAL": "Piso (TOTAL)", "ganancia_TOTAL": "Ganancia sobre el piso (TOTAL)",
    "piso_media": "Piso (F1 contra otras entrevistas)", "ganancia_media": "Ganancia sobre el piso",
    "pct_ganancia_positiva": "% entrevistas con ganancia > 0",
    "rango_medio": "Rango medio (1 = la propia es la más parecida)",
    "pct_rango_1": "% de casos en que la propia es la más parecida",
    "candidatos": "Entrevistas comparadas", "medida": "Medida",
    "parecido_humano": "Parecido entre entrevistas: humano",
    "parecido_modelo": "Parecido entre entrevistas: modelo",
    "dif_parecido": "Diferencia (modelo − humano)",
    "preocupacion": "Preocupación", "causa": "Causas", "consecuencia": "Consecuencias",
    "accion": "Acciones", "TOTAL": "TOTAL",
}


def _write_table(ws, df: pd.DataFrame, row0: int, title: Optional[str] = None,
                 heat_cols: tuple = (), note: Optional[str] = None, heat_relativo: bool = False) -> int:
    from openpyxl.formatting.rule import ColorScaleRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    r = row0
    if title:
        ws.cell(r, 1, title).font = Font(bold=True, size=12, color="1F3864")
        r += 1
    if note:
        ws.cell(r, 1, note).font = Font(italic=True, size=9, color="606060")
        r += 1
    thin = Side(style="thin", color="C8C8C8")
    brd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for j, col in enumerate(df.columns, 1):
        c = ws.cell(r, j, ENCABEZADOS.get(col, col))
        c.fill = PatternFill("solid", fgColor="1F3864")
        c.font = Font(color="FFFFFF", bold=True)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = brd
    ws.row_dimensions[r].height = 32
    hdr = r
    for i, rec in enumerate(df.itertuples(index=False), 1):
        for j, v in enumerate(rec, 1):
            if isinstance(v, float) and math.isnan(v):
                v = None
            col = df.columns[j - 1]
            if col == "campo":
                v = ENCABEZADOS.get(v, v)
            elif col == "metrica":
                v = METRICA_NOMBRE.get(v, v)
            c = ws.cell(hdr + i, j, v)
            c.border = brd
            if isinstance(v, float):
                c.number_format = "0.0#" if df.columns[j - 1] == "temperatura" else "0.0000"
                c.alignment = Alignment(horizontal="center")
    last = hdr + len(df)
    if len(df):
        low, mid, high = _HEATMAP_THRESH
        for j, col in enumerate(df.columns, 1):
            if col in heat_cols:
                L = get_column_letter(j)
                if heat_relativo:
                    # Escala relativa a la propia columna: mínimo rojo, mediana amarilla, máximo verde
                    regla = ColorScaleRule(
                        start_type="min", start_color=_HEATMAP[0],
                        mid_type="percentile", mid_value=50, mid_color=_HEATMAP[1],
                        end_type="max", end_color=_HEATMAP[2])
                else:
                    regla = ColorScaleRule(
                        start_type="num", start_value=low, start_color=_HEATMAP[0],
                        mid_type="num", mid_value=mid, mid_color=_HEATMAP[1],
                        end_type="num", end_value=high, end_color=_HEATMAP[2])
                ws.conditional_formatting.add(f"{L}{hdr + 1}:{L}{last}", regla)
    for j, col in enumerate(df.columns, 1):
        L = get_column_letter(j)
        largo = max([len(str(ENCABEZADOS.get(col, col)))] +
                    [len(str(v)) for v in df[col].head(200)]) if len(df) else len(str(col))
        ws.column_dimensions[L].width = max(ws.column_dimensions[L].width or 0, min(max(10, largo + 2), 60))
    return last + 2


def _metodologia(res: dict, human: HumanCoding, metrics: list[str], bertscore_model: str,
                 referencia: str) -> list[tuple[str, str]]:
    conds = res["conds"]
    reps = ", ".join(f"{c.nombre}: {len(c.runs)}" for c in conds)
    return [
        ("Qué es este archivo",
         "Compara las etiquetas del pipeline contra la codificación humana oficial. Cada fila de "
         "'Resumen' es una condición experimental (experimento + prompt)."),
        ("Repeticiones",
         "El modelo no responde exactamente igual en corridas idénticas, aun con temperatura 0. "
         "Por eso cada condición se corre varias veces (ej. E1, E1_rep2, E1_rep3) y aquí se "
         f"agrupan como una sola condición. Repeticiones por condición: {reps}."),
        ("Cómo se combinan las repeticiones",
         "Las repeticiones NO se juntan como una muestra más grande (serían las mismas entrevistas "
         "varias veces y se inflaría la n). Para cada entrevista y campo se promedia el F1 de sus "
         "repeticiones; queda una calificación por entrevista y sobre esas se calculan los "
         "estadísticos y la comparación pareada. La n es el número de entrevistas."),
        ("Ruido entre repeticiones",
         "La hoja 'Ruido' muestra la media de cada repetición por separado, su rango y cuánto varía "
         "cada entrevista entre repeticiones. Es la línea base: si la diferencia entre dos "
         "condiciones es del tamaño de ese ruido, no es una diferencia real."),
        ("Carpetas",
         "Cada repetición conserva su propia carpeta en outputs_v<versión>/ y su propia evaluación "
         f"en {EVAL_DIRNAME}/. Solo este comparativo las junta al leerlas."),
        ("Métrica",
         f"Métricas calculadas: {', '.join(metrics)}. Métrica principal (hoja 'Resumen'): "
         f"{res['principal']}." + (f" Modelo de BERTScore: {bertscore_model}." if "bertscore" in metrics else "")),
        ("Calificación de un campo",
         "Recall: para cada etiqueta humana se toma su mejor pareja entre las del modelo y se "
         "promedian (¿cuánto de lo humano capturó el modelo?). Precisión: lo mismo desde las "
         "etiquetas del modelo (¿cuánto de lo del modelo tiene respaldo humano?). F1: media armónica."),
        ("TOTAL",
         "Para cada entrevista, promedio de los campos que tienen F1; luego estadísticos entre entrevistas."),
        ("Piso aleatorio",
         "Todas las entrevistas hablan de agua en la CDMX, así que dos etiquetas cualesquiera del "
         "tema ya se parecen algo y BERTScore les da una calificación alta aunque no correspondan. "
         "Para saber cuánto de la calificación es captar ESA entrevista y cuánto es solo hablar del "
         "mismo tema, las etiquetas del modelo de cada entrevista se califican también contra la "
         "codificación humana de cada una de las demás entrevistas. El promedio de esas "
         "calificaciones cruzadas es el piso. Ganancia = F1 propio − piso: si es cercana a 0, el "
         "modelo no distingue esa entrevista de las demás. No se sortea: se usan todas las demás "
         "entrevistas, así que el resultado no depende del azar. Rango: lugar que ocupa la propia "
         "codificación humana entre todas (1 = la más parecida). Un prompt que produzca etiquetas "
         "genéricas puede subir el F1 y también el piso; por eso la comparación pareada se hace "
         "sobre F1 y sobre la ganancia."),
        ("Parecido entre entrevistas",
         "Las entrevistas comparten un tema y es esperable que sus narrativas se traslapen; lo que "
         "interesa es si el modelo reproduce ese traslape o lo exagera. Para cada campo se califica "
         "la codificación humana de cada entrevista contra la codificación humana de cada una de las "
         "otras (parecido humano) y lo mismo con las etiquetas del modelo (parecido modelo), con la "
         "misma fórmula de F1 y solo entre entrevistas que tienen F1 en ese campo. Diferencia = "
         "modelo − humano: cerca de 0, el modelo reproduce cuánto se parecen las narrativas; positiva, "
         "las aplana (hace que las entrevistas se parezcan más entre sí); negativa, exagera sus "
         "diferencias. Ojo: la codificación humana la hicieron varias personas con estilos distintos "
         "y el modelo es un solo 'codificador', así que parte de la diferencia puede ser de estilo "
         "(más en chrF y BLEU, que comparan palabras, que en BERTScore, que compara significado)."),
        ("Campos vacíos",
         "El F1 se calcula solo donde humano y modelo tienen etiquetas. Los demás casos (ambos vacíos, "
         "solo humano, solo modelo) se cuentan en la hoja 'Presencia' y no entran al F1."),
        ("Comparación pareada",
         f"Cada condición se compara entrevista por entrevista contra la referencia ({referencia}) y "
         "contra la condición anterior con el mismo prompt (si no hay, la del experimento previo). "
         "Se reporta la diferencia media y mediana, cuántas entrevistas mejoran o empeoran y la "
         "prueba de Wilcoxon de rangos con signo (datos pareados)."),
        ("Codificación humana",
         f"{human.path.name}, hoja '{human.sheet}' (sha256 {human.sha256[:12]}…). La hoja "
         "'Referencia humana' muestra cómo se leyó cada celda."),
        ("Reglas de lectura", " · ".join(REGLAS)),
        ("Colores", f"Mismo semáforo con umbrales fijos para todas las métricas: rojo ≤ "
                    f"{_HEATMAP_THRESH[0]}, amarillo ≈ {_HEATMAP_THRESH[1]}, verde ≥ {_HEATMAP_THRESH[2]}. "
                    "chrF++ y BLEU comparan palabras literales y suelen salir en rojo: indica que hay "
                    "poca coincidencia literal con la codificación humana."),
        ("Versión", f"Evaluación v{EVAL_VERSION} ({Path(__file__).name})."),
    ]


def export_excel(res: dict, human: HumanCoding, path: Path, metrics: list[str],
                 bertscore_model: str, referencia: str) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    t, principal = res["tablas"], res["principal"]
    wb = Workbook()
    wb.remove(wb.active)

    # Una hoja de resumen por métrica; la principal primero. Todas usan el mismo semáforo de
    # umbrales fijos, para que el color diga lo mismo en todas (rojo = poca coincidencia).
    cols_piso = ["piso_media", "ganancia_media", "pct_ganancia_positiva", "rango_medio",
                 "pct_rango_1", "candidatos"]
    cols_par = ["parecido_humano", "parecido_modelo", "dif_parecido"]
    for met in metrics:
        ws = wb.create_sheet("Resumen" if met == principal else f"Resumen {METRICA_NOMBRE[met]}")
        rel = False
        aviso = ""
        titulo = f"MEGADAPT — Evaluación v{EVAL_VERSION} · {METRICA_NOMBRE[met]} F1"
        if met == principal:
            titulo += " (métrica principal)"
        ws.cell(1, 1, titulo).font = Font(bold=True, size=13, color="1F3864")
        cond = t["condiciones"]
        cond = cond[cond["metrica"] == met].drop(columns=["condicion", "metrica"])
        r = _write_table(ws, cond, 3, "F1 promedio por condición",
                         heat_cols=tuple(CAMPOS + ["TOTAL"]), heat_relativo=rel,
                         note="Una fila por condición. Las repeticiones se promedian por entrevista "
                              "(ver 'Metodología')." + aviso)
        dist = t["distribucion"]
        dist = dist[dist["metrica"] == met].drop(columns=["condicion", "metrica"])
        r = _write_table(ws, dist.drop(columns=cols_piso + cols_par), r, "Distribución entre entrevistas",
                         heat_cols=("media", "min", "q1", "mediana", "q3", "max"), heat_relativo=rel,
                         note="n = entrevistas con F1 en ese campo. Etiquetas: promedio por entrevista." + aviso)
        piso = dist[["experimento", "prompt", "campo", "media"] + cols_piso]
        r = _write_table(ws, piso, r, "Piso aleatorio: ¿se parece más a su entrevista que a las demás?",
                         heat_cols=("media", "piso_media"), heat_relativo=rel,
                         note="Piso = F1 de las etiquetas del modelo contra la codificación humana de las "
                              "OTRAS entrevistas. Ganancia = F1 − piso (ver 'Metodología')." + aviso)
        par = dist[["experimento", "prompt", "campo"] + cols_par]
        _write_table(ws, par, r, "¿El modelo aplana las diferencias entre entrevistas?",
                     note="Parecido = F1 entre entrevistas distintas (humano contra humano; modelo contra "
                          "modelo). Diferencia > 0: el modelo hace que las entrevistas se parezcan más "
                          "de lo que se parecen según el humano.")
        ws.freeze_panes = "C6"

    w = wb.create_sheet("Metodología")
    w.column_dimensions["A"].width = 30
    w.column_dimensions["B"].width = 110
    w.cell(1, 1, "Metodología").font = Font(bold=True, size=13, color="1F3864")
    for i, (k, v) in enumerate(_metodologia(res, human, metrics, bertscore_model, referencia), 3):
        w.cell(i, 1, k).font = Font(bold=True)
        c = w.cell(i, 2, v)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        w.cell(i, 1).alignment = Alignment(vertical="top")

    hojas = [
        ("Ruido", t["ruido"], "Variación entre repeticiones de una misma condición", ()),
        ("Comparación pareada", t["pareado"],
         "Diferencia por entrevista (condición − contra), en F1 y en ganancia sobre el piso; "
         "positiva = mejora", ()),
        ("Pareado detalle", t["pareado_detalle"],
         f"Diferencia de cada entrevista ({METRICA_NOMBRE[principal]} F1)", ("f1_condicion", "f1_contra")),
        ("Presencia", t["presencia"], "Campos con y sin etiquetas (suma de repeticiones)", ()),
        ("Corridas", t["corridas"], "Corridas incluidas en este comparativo", ()),
    ]
    for nombre, df, titulo, heat in hojas:
        w = wb.create_sheet(nombre)
        _write_table(w, df, 1, titulo, heat_cols=heat)
        w.freeze_panes = "A3"

    # Por entrevista: una tabla por métrica en la misma hoja
    w = wb.create_sheet("Por entrevista")
    rr = 1
    conds_cols = tuple(c.nombre for c in res["conds"])
    for met in metrics:
        sub = t["por_entrevista"]
        sub = sub[sub["metrica"] == met].drop(columns="metrica")
        rr = _write_table(w, sub, rr, f"{METRICA_NOMBRE[met]} F1 por entrevista y campo (promedio de repeticiones)",
                          heat_cols=conds_cols)

    w = wb.create_sheet("Referencia humana")
    rr = 1
    if human.filas_omitidas:
        om = pd.DataFrame(human.filas_omitidas, columns=["Fila de Excel", "Sector"])
        rr = _write_table(w, om, rr, "Filas de bloque omitidas (no pertenecen a ninguna entrevista)")
    emparejadas = {h for r in res["conds"] for run in r.runs for h in run.df["Id_humano"].unique()}
    ref = pd.DataFrame([{
        "Clave": k, "Emparejada": "sí" if k in emparejadas else "no", "campo": c,
        "N etiquetas": len(human.data[k][c]), "Etiquetas": " | ".join(human.data[k][c]),
        "Encabezados": " | ".join(human.heads[k][c]),
    } for k in sorted(human.data) for c in CAMPOS])
    _write_table(w, ref, rr, "Cómo se leyó cada celda del Excel humano")
    wb.save(path)


def write_comparativo(runs: list[Run], human: HumanCoding, metrics: list[str], bertscore_model: str,
                      referencia: str, destino: Optional[str]) -> Path:
    firma = check_same_evaluation(runs)
    res = build_comparativo(runs, metrics, referencia)
    if destino:
        base = Path(destino).resolve()
    else:
        roots = {_outputs_root(r.carpeta) for r in runs}
        if len(roots) != 1 or None in roots:
            raise SystemExit("Las corridas no están en una sola carpeta outputs_v<versión>/. "
                             "Indica dónde guardar el comparativo con --comparativo-en.")
        base = roots.pop() / "comparativos"
    out = base / datetime.now().strftime("%Y%m%d_%H%M%S")
    n = 2
    while out.exists():
        out = out.with_name(f"{out.name}_{n}")
        n += 1
    out.mkdir(parents=True)
    try:
        for nombre, df in res["tablas"].items():
            df.to_csv(out / f"{nombre}.csv", index=False, encoding="utf-8-sig")
        export_excel(res, human, out / "megadapt_evaluacion.xlsx", metrics, bertscore_model, referencia)
    except Exception:
        shutil.rmtree(out, ignore_errors=True)   # no dejar comparativos a medias
        raise
    cfg = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "script": Path(__file__).name,
        "eval_version": EVAL_VERSION,
        "firma_evaluacion": firma,
        "excel_humano": str(human.path.resolve()),
        "metrica_principal": res["principal"],
        "referencia": referencia,
        "condiciones": [{"condicion": c.nombre, "corridas": [r.nombre for r in c.runs]} for c in res["conds"]],
        "comparaciones": [{"condicion": c.nombre, "contra": b.nombre, "tipo": t} for c, b, t in res["comparaciones"]],
    }
    (out / "comparativo_config.json").write_text(json.dumps(cfg, ensure_ascii=False, indent=2), encoding="utf-8")

    for met in metrics:
        nombre = METRICA_NOMBRE[met] + (" · métrica principal" if met == res["principal"] else "")
        print("\n" + "=" * 78)
        print(f"{nombre}")
        print("=" * 78)
        cond = res["tablas"]["condiciones"]
        print("F1 PROMEDIO POR CONDICIÓN")
        print(cond[cond["metrica"] == met].drop(columns=["condicion", "corridas", "metrica"]).to_string(index=False))
        dist = res["tablas"]["distribucion"]
        dist = dist[dist["metrica"] == met]
        print("\nPISO ALEATORIO (¿se parece más a su entrevista que a las demás?)")
        print(dist[["condicion", "campo", "media", "piso_media", "ganancia_media", "pct_ganancia_positiva",
                    "rango_medio", "pct_rango_1", "candidatos"]].to_string(index=False))
        print("\nPARECIDO ENTRE ENTREVISTAS (diferencia > 0: el modelo aplana las diferencias)")
        print(dist[["condicion", "campo", "parecido_humano", "parecido_modelo", "dif_parecido"]].to_string(index=False))
        par = res["tablas"]["pareado"]
        par = par[(par["metrica"] == met) & (par["campo"] == "TOTAL")]
        if len(par):
            print("\nCOMPARACIÓN PAREADA (TOTAL)")
            print(par[["condicion", "contra", "tipo", "medida", "n", "diferencia_media", "mejoran", "empeoran",
                       "p_wilcoxon", "ruido_rango_contra"]].to_string(index=False))
    print(f"\nComparativo: {out}")
    return out


# ──────────────────────────────────────────────────────────────────────────────
# 8. CLI
# ──────────────────────────────────────────────────────────────────────────────
def main(argv=None):
    ap = argparse.ArgumentParser(
        description=f"Evaluación NLP MEGADAPT v{EVAL_VERSION}: corridas del pipeline vs. codificación humana",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=f"""
Ejemplos (cada comando va en una sola línea):

  # Evaluar todas las corridas de la versión 7.3 y armar el comparativo:
  python megadapt_eval_nlp_v4.py --runs outputs_v7.3

  # Solo algunas corridas (carpetas dentro de outputs_v7.3):
  python megadapt_eval_nlp_v4.py --runs outputs_v7.3/<corrida_E1> outputs_v7.3/<corrida_E1_rep2>

  # Agregar chrF y BLEU además de BERTScore:
  python megadapt_eval_nlp_v4.py --runs outputs_v7.3 --metrics bertscore chrf bleu

Salidas:
  <corrida>/{EVAL_DIRNAME}/            eval_config.json, eval_por_entrevista.csv, eval_resumen.csv
  outputs_v<versión>/comparativos/<fecha_hora>/   megadapt_evaluacion.xlsx, CSVs y comparativo_config.json

El Excel humano se toma de --human o de MEGADAPT_HUMANO en el archivo .env.
""",
    )
    ap.add_argument("--runs", nargs="+", required=True,
                    help="Carpetas de corrida, o una carpeta que las contenga (ej. outputs_v7.3).")
    ap.add_argument("--human", default=os.environ.get("MEGADAPT_HUMANO"),
                    help="Excel de codificación humana. Si no se pone, se usa MEGADAPT_HUMANO del .env.")
    ap.add_argument("--sheet", default="Integracion", help="Hoja del Excel humano (default: Integracion).")
    ap.add_argument("--metrics", nargs="+", default=["bertscore"], choices=list(METRICAS),
                    help="Métricas (default: bertscore). La principal es bertscore > chrf > bleu.")
    ap.add_argument("--bertscore-model", default=BERTSCORE_MODEL_DEFAULT,
                    help=f"Modelo de BERTScore (default y fijo para el proyecto: {BERTSCORE_MODEL_DEFAULT}).")
    ap.add_argument("--device", default="cpu", help="Dispositivo para BERTScore: cpu o cuda.")
    ap.add_argument("--referencia", default="E1", help="Experimento de referencia para la comparación (default: E1).")
    ap.add_argument("--reevaluar", action="store_true",
                    help=f"Reemplaza evaluaciones {EVAL_DIRNAME}/ hechas con otra configuración.")
    ap.add_argument("--sin-comparativo", action="store_true",
                    help="Solo evalúa cada corrida; no arma el comparativo.")
    ap.add_argument("--comparativo-en", default=None,
                    help="Carpeta donde guardar el comparativo (por defecto outputs_v<versión>/comparativos/).")
    a = ap.parse_args(argv)

    if not a.human:
        ap.error("Falta --human (o MEGADAPT_HUMANO en el archivo .env).")
    hpath = Path(a.human)
    if not hpath.exists():
        ap.error(f"No existe el Excel humano: {hpath}")
    if a.bertscore_model != BERTSCORE_MODEL_DEFAULT and "bertscore" in a.metrics:
        log.warning(f"Usas {a.bertscore_model} en lugar del modelo fijo del proyecto ({BERTSCORE_MODEL_DEFAULT}).")
    metrics = [m for m in METRICAS if m in a.metrics]

    human = load_human(hpath, a.sheet)
    log.info(f"Codificación humana: {len(human.data)} claves leídas de {hpath.name} "
             f"({len(human.filas_omitidas)} filas de bloque omitidas).")
    runs = find_runs(a.runs)
    log.info(f"Corridas encontradas: {len(runs)}")

    sims = {}

    def get_sims():
        for m in metrics:
            if m not in sims:
                sims[m] = make_sim(m, a.bertscore_model, a.device)
        return sims

    for r in runs:
        evaluate_or_load(r, human, metrics, a.bertscore_model, get_sims, a.reevaluar)

    if not a.sin_comparativo:
        write_comparativo(runs, human, metrics, a.bertscore_model, a.referencia, a.comparativo_en)


if __name__ == "__main__":
    main()
