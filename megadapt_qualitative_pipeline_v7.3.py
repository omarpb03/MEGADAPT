"""
MEGADAPT - Pipeline de análisis cualitativo automatizado, v7.3 (OpenAI)
Framework de experimentación con variantes de prompt

Cambios respecto a v7.2:
  A. Prompts solo desde la carpeta PROMPTS/: un archivo .md completo por variante.
     Ya no hay prompts escritos en el código ni respaldo silencioso: si falta la
     carpeta o un .md, el script se detiene.
  B. El texto del .md llega al modelo tal cual (ya no se rellena {max_items}).
     Si un .md contiene {max_items}, el script marca error.
  C. --max-items ya no va en el prompt: es un tope de seguridad que se aplica en
     el postproceso, después de limpiar y deduplicar. Por defecto no recorta.
     Cada recorte queda registrado en el JSON individual (items_antes_de_tope),
     en la terminal y en usage_log_v7.jsonl (entrevistas_recortadas).
  D. --run-all-variants revisa los .md de todas las variantes antes de llamar a la API.
  E. Cada corrida se guarda en su propia carpeta dentro de outputs_v<versión>/, junto
     al script (<fecha_hora>_<variante>_<modelo>/). Nada se sobrescribe. --out solo
     sirve para retomar una corrida cortada, y se detiene si la configuración cambió.
  F.   F. El preproceso ya no renombra E:/R: ni borra lo que está entre paréntesis o
     corchetes. Las etiquetas no son sistemáticas (en algunas transcripciones "E"
     marca al entrevistado) y dentro de ( ) y [ ] hay etiquetas de hablante, nombres
     de instituciones y contenido sustantivo. Las transcripciones se mandan tal cual.
  G. --experimento (etiqueta, ej. E1) y --temperatura (default 0.0). Las dos quedan en
     experiment_config.json, y la etiqueta va también en el nombre de la carpeta
     (<fecha_hora>_<experimento>_<variante>_<modelo>/). La temperatura registrada es
  H. Antes de crear la carpeta, verifica la llave de OpenAI con una llamada sin costo.
     Si es inválida, expiró o no hay conexión, se detiene sin dejar rastro.
  I. --esquema: formulario de salida. con_largo (default, el de E1: descripciones con
     "4 o 5 palabras") o sin_largo (mismas claves, sin número de palabras). Queda en
     experiment_config.json junto con el formulario completo (esquema_json).
  J. La carpeta de transcripciones puede venir de MEGADAPT_TRANSCRIPCIONES en el .env;
     queda anotada en experiment_config.json (carpeta_transcripciones).

Variantes de prompt (diseño aditivo)
  p1_baseline      Prompt baseline diseñado por el equipo.
  p2_definiciones  Baseline + definiciones de etiquetado.
  p3_perfil        Baseline + definiciones + descripción del perfil cualitativo.

Se mantiene de v7: limpieza de etiquetas nulas (bug "Original"), checkpoints por
entrevista y modo --count-only.

Uso: ver ejemplos con --help.

Requiere OPENAI_API_KEY en .env (salvo con --count-only).
"""
from __future__ import annotations

import argparse
import json
import logging
import os
import re
import threading
from concurrent.futures import ThreadPoolExecutor, as_completed
from dataclasses import asdict, dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Optional

import pandas as pd
from dotenv import load_dotenv
from pydantic import BaseModel, Field
from tenacity import (before_sleep_log, retry, retry_if_exception,
                      stop_after_attempt, wait_exponential)
from tqdm import tqdm

load_dotenv()

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ──────────────────────────────────────────────────────────────────────────────
# 0. Catálogo de modelos
# ──────────────────────────────────────────────────────────────────────────────


@dataclass(frozen=True)
class ModelInfo:
    context: int
    price_in: float   # USD por 1M tokens de entrada
    price_out: float  # USD por 1M tokens de salida


MODEL_CATALOG: dict[str, ModelInfo] = {
    "gpt-4o":        ModelInfo(128_000,   2.50, 10.00),
    "gpt-4o-mini":   ModelInfo(128_000,   0.15,  0.60),
    "gpt-4.1":       ModelInfo(1_047_576, 2.00,  8.00),
    "gpt-4.1-mini":  ModelInfo(1_047_576, 0.40,  1.60),
    "gpt-4.1-nano":  ModelInfo(1_047_576, 0.10,  0.40),
}

TEMPERATURE_PREFIXES = ("gpt-4o", "gpt-4.1")
PROMPT_OVERHEAD_TOKENS = 400
EST_OUTPUT_TOKENS = 900
CONTEXT_SAFETY = 0.90

# Versión del script: define la carpeta de resultados (outputs_v<versión>/)
# y queda anotada en cada experiment_config.json. Cámbiala al crear una versión nueva.
SCRIPT_VERSION = "7.3"
OUTPUTS_ROOT = Path(__file__).parent / f"outputs_v{SCRIPT_VERSION}"

# ──────────────────────────────────────────────────────────────────────────────
# 1. Estructuras Pydantic y dataclasses
# ──────────────────────────────────────────────────────────────────────────────


@dataclass
class InterviewMetadata:
    Id_entrevista: str


@dataclass
class InterviewResult:
    metadata: InterviewMetadata
    prompt_variant: str = "p1_baseline"
    Preocupacion_principal_acerca_del_agua: list[str] = field(
        default_factory=list)
    Causas_principales: list[str] = field(default_factory=list)
    Consecuencias: list[str] = field(default_factory=list)
    Acciones_y_Actores_que_realizan_dichas_acciones: list[dict] = field(
        default_factory=list)
    input_tokens: int = 0
    output_tokens: int = 0
    error: Optional[str] = None
    # {campo: n_original} solo si se recortó
    items_antes_de_tope: dict = field(default_factory=dict)


class AccionActor(BaseModel):
    actor: str = Field(
        description="Quién realiza la acción (institución, grupo o persona), en 1 a 3 palabras."
    )
    accion: str = Field(
        description="La acción, resumida en 4 o 5 palabras."
    )


class QualitativeAnalysisSchema(BaseModel):
    Preocupacion_principal_acerca_del_agua: list[str] = Field(
        description="Preocupaciones principales sobre el agua, cada una de 4 o 5 palabras."
    )
    Causas_principales: list[str] = Field(
        description="Causas del problema mencionadas, cada una de 4 o 5 palabras."
    )
    Consecuencias: list[str] = Field(
        description="Consecuencias o efectos mencionados, cada una de 4 o 5 palabras."
    )
    Acciones_y_Actores_que_realizan_dichas_acciones: list[AccionActor] = Field(
        description="Acciones realizadas o propuestas, con el actor que las realiza."
    )


# Esquema sin largo fijo: mismas claves, pero las descripciones ya no piden un número de palabras.
class AccionActorSinLargo(BaseModel):
    actor: str = Field(
        description="Quién realiza la acción (institución, grupo o persona).")
    accion: str = Field(
        description="La acción que ese actor realiza o propone.")


class QualitativeAnalysisSchemaSinLargo(BaseModel):
    Preocupacion_principal_acerca_del_agua: list[str] = Field(
        description="Preocupaciones principales acerca del agua mencionadas en la entrevista."
    )
    Causas_principales: list[str] = Field(
        description="Causas principales del problema mencionadas en la entrevista."
    )
    Consecuencias: list[str] = Field(
        description="Consecuencias o efectos mencionados en la entrevista."
    )
    Acciones_y_Actores_que_realizan_dichas_acciones: list[AccionActorSinLargo] = Field(
        description="Acciones realizadas o propuestas, con el actor que las realiza."
    )


# Esquemas disponibles para --esquema. "con_largo" es el que se usó en E1.
ESQUEMAS = {
    "con_largo": QualitativeAnalysisSchema,
    "sin_largo": QualitativeAnalysisSchemaSinLargo,
}


# ──────────────────────────────────────────────────────────────────────────────
# 2. Variantes de prompt
# ──────────────────────────────────────────────────────────────────────────────
# El texto de cada prompt vive SOLO en PROMPTS/<nombre>.md (un archivo completo
# por variante). Aquí solo se declaran los nombres y el orden de --run-all-variants.
PROMPT_VARIANTS: list[str] = [
    "p1_baseline",
    "p2_definiciones",
    "p3_perfil",
]

USER_TEMPLATE = "=== TRANSCRIPCIÓN ===\n{transcript}\n=== FIN DE TRANSCRIPCIÓN ==="

# Palabras que indican que el LLM no encontró contenido real — se eliminan en postproceso
_NULL_LABELS = {
    "original", "none", "ninguno", "ninguna", "na", "n/a", "n a",
    "no aplica", "no especificado como accion", "acción no especificada",
    "no hay accion", "sin accion", "acción desconocida",
}

# ──────────────────────────────────────────────────────────────────────────────
# 3. Lectura y preprocesamiento
# ──────────────────────────────────────────────────────────────────────────────


def load_transcript(path: Path) -> str:
    if path.suffix.lower() == ".docx":
        from docx import Document
        return "\n".join(p.text.strip() for p in Document(str(path)).paragraphs if p.text.strip())
    if path.suffix.lower() in (".txt", ".md"):
        return path.read_text(encoding="utf-8", errors="replace")
    return ""


def parse_interview_id(filename: str) -> InterviewMetadata:
    stem = Path(filename).stem.upper()
    base, _, transcript_label = stem.partition("__T-")
    interview_id = base
    if "PARTE" in transcript_label:
        part_label = re.search(
            r"(\d+RA|\d+DA|\d+TA)\s*PARTE", transcript_label)
        if part_label:
            interview_id = f"{base}_{part_label.group(0).replace(' ', '_')}"
    return InterviewMetadata(Id_entrevista=interview_id)


def preprocess(text: str) -> str:
    """Limpieza ligera; no trunca."""
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text


# ──────────────────────────────────────────────────────────────────────────────
# 4. Tokens y costos
# ──────────────────────────────────────────────────────────────────────────────
_ENCODER = None
TOKEN_MODE = "sin inicializar"


def _get_encoder():
    global _ENCODER, TOKEN_MODE
    if _ENCODER is None:
        try:
            import tiktoken
            _ENCODER = tiktoken.get_encoding("o200k_base")
            TOKEN_MODE = "exacto (tiktoken o200k_base)"
        except Exception as exc:
            log.warning(
                f"tiktoken no disponible ({type(exc).__name__}); aproximación 1.5 tok/pal.")
            _ENCODER = False
            TOKEN_MODE = "APROXIMADO (1.5 tokens/palabra)"
    return _ENCODER


def count_tokens(text: str) -> int:
    enc = _get_encoder()
    return len(enc.encode(text, disallowed_special=())) if enc else int(len(text.split()) * 1.5)


def estimate_cost(model: str, input_tokens: int, output_tokens: int) -> Optional[float]:
    info = MODEL_CATALOG.get(model)
    if not info:
        return None
    return input_tokens * info.price_in / 1e6 + output_tokens * info.price_out / 1e6


def list_transcripts(folder: str) -> list[Path]:
    return sorted(
        f for f in Path(folder).rglob("*")
        if not f.name.startswith("~$") and f.suffix.lower() in (".docx", ".txt", ".md")
    )


def count_only(folder: str, out_dir: str, models: list[str], variant: str = "p1_baseline") -> pd.DataFrame:
    files = list_transcripts(folder)
    if not files:
        raise SystemExit(f"No se encontraron transcripciones en {folder}")
    prompt = build_prompt(variant)
    prompt_tokens = count_tokens(prompt) + PROMPT_OVERHEAD_TOKENS
    rows = []
    for f in tqdm(files, desc="Contando tokens"):
        text = preprocess(load_transcript(f))
        toks = count_tokens(text)
        row = {
            "Id_entrevista": parse_interview_id(f.name).Id_entrevista,
            "archivo": f.name,
            "palabras": len(text.split()),
            "tokens_transcripcion": toks,
            "tokens_entrada_total": toks + prompt_tokens,
        }
        for m in models:
            info = MODEL_CATALOG[m]
            row[f"pct_contexto_{m}"] = round(
                100 * (toks + prompt_tokens) / info.context, 1)
            row[f"costo_usd_{m}"] = round(estimate_cost(
                m, toks + prompt_tokens, EST_OUTPUT_TOKENS), 4)
        rows.append(row)
    df = pd.DataFrame(rows)
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    df.to_csv(out / "token_report.csv", index=False, encoding="utf-8-sig")

    summary = []
    for m in models:
        info = MODEL_CATALOG[m]
        exceeds = int((df["tokens_entrada_total"] >
                      CONTEXT_SAFETY * info.context).sum())
        summary.append({
            "modelo": m,
            "context_window": info.context,
            "precio_in_por_1M": info.price_in,
            "precio_out_por_1M": info.price_out,
            "max_pct_contexto": df[f"pct_contexto_{m}"].max(),
            "entrevistas_exceden_90pct": exceeds,
            "costo_total_usd_estimado": round(df[f"costo_usd_{m}"].sum(), 2),
        })
    sdf = pd.DataFrame(summary)
    sdf.to_csv(out / "model_cost_summary.csv",
               index=False, encoding="utf-8-sig")
    print(f"\nModo de conteo: {TOKEN_MODE}")
    print(f"Prompt variant: {variant} ({prompt_tokens} tokens de overhead)")
    print(
        f"Entrevistas: {len(df)} | palabras totales: {df['palabras'].sum():,}")
    print(sdf.to_string(index=False))
    print("\nPrecios de referencia: verificar en openai.com antes de correr.")
    return df


# ──────────────────────────────────────────────────────────────────────────────
# 5. Construcción del prompt — carga desde la carpeta PROMPTS
# ──────────────────────────────────────────────────────────────────────────────
_ACTIVE_PROMPTS_DIR: Optional[Path] = None
_DEFAULT_PROMPTS_DIR = Path(__file__).parent / "PROMPTS"


def set_prompts_dir(path: Optional[str]) -> None:
    """
    Fija la carpeta de la que se leen los prompts (.md).
    Sin --prompts-dir se usa PROMPTS/ junto al script. Si no existe, el script
    se detiene: no hay prompts de respaldo en el código.
    """
    global _ACTIVE_PROMPTS_DIR
    p = Path(path) if path is not None else _DEFAULT_PROMPTS_DIR
    if not p.is_dir():
        raise FileNotFoundError(
            f"Carpeta de prompts no encontrada: {p.resolve()}\n"
            f"Debe contener un .md por variante: {', '.join(PROMPT_VARIANTS)}"
        )
    _ACTIVE_PROMPTS_DIR = p
    log.info(f"Prompts desde: {p.resolve()}")


def build_prompt(variant: str) -> str:
    """
    Devuelve el texto de PROMPTS/<variant>.md tal cual, sin rellenar nada.
    Lo que está en el archivo es exactamente lo que recibe el modelo.
    """
    if variant not in PROMPT_VARIANTS:
        raise ValueError(
            f"Variante '{variant}' no existe. Opciones: {PROMPT_VARIANTS}")
    if _ACTIVE_PROMPTS_DIR is None:
        raise RuntimeError(
            "Carpeta de prompts sin configurar (falta set_prompts_dir).")
    md_path = _ACTIVE_PROMPTS_DIR / f"{variant}.md"
    if not md_path.exists():
        raise FileNotFoundError(
            f"Falta el prompt de '{variant}': {md_path.resolve()}")
    text = md_path.read_text(encoding="utf-8-sig").strip()
    if not text:
        raise ValueError(f"El prompt {md_path.name} está vacío.")
    if "{max_items}" in text:
        raise ValueError(
            f"{md_path.name} contiene {{max_items}}. El límite de ítems va solo "
            f"en el postproceso (--max-items); quítalo del prompt."
        )
    return text

# ──────────────────────────────────────────────────────────────────────────────
# 6. Cliente OpenAI con reintentos
# ──────────────────────────────────────────────────────────────────────────────


class UsageTracker:
    def __init__(self, model: str):
        self.model, self.input, self.output = model, 0, 0
        self._lock = threading.Lock()

    def add(self, usage) -> tuple[int, int]:
        i = getattr(usage, "prompt_tokens", 0) or 0
        o = getattr(usage, "completion_tokens", 0) or 0
        with self._lock:
            self.input += i
            self.output += o
        return i, o

    def cost(self) -> Optional[float]:
        return estimate_cost(self.model, self.input, self.output)


def _is_retryable(exc: BaseException) -> bool:
    import openai
    if getattr(exc, "code", None) == "insufficient_quota":
        return False
    if isinstance(exc, (openai.RateLimitError, openai.APIConnectionError,
                        openai.APITimeoutError, openai.InternalServerError)):
        return True
    msg = str(exc).lower()
    return any(k in msg for k in ("503", "429", "rate limit", "overloaded"))


def _call_structured(client, model: str, system: str, user: str, schema, temperature: Optional[float]):
    @retry(
        retry=retry_if_exception(_is_retryable),
        wait=wait_exponential(multiplier=1, min=2, max=60),
        stop=stop_after_attempt(6),
        before_sleep=before_sleep_log(log, logging.WARNING),
        reraise=True,
    )
    def _do():
        kwargs = dict(
            model=model,
            messages=[
                {"role": "system", "content": system},
                {"role": "user",   "content": user},
            ],
            response_format=schema,
        )
        if temperature is not None and model.startswith(TEMPERATURE_PREFIXES):
            kwargs["temperature"] = temperature
        resp = client.chat.completions.parse(**kwargs)
        msg = resp.choices[0].message
        if getattr(msg, "refusal", None):
            raise RuntimeError(
                f"El modelo rechazó la solicitud: {msg.refusal}")
        if msg.parsed is None:
            raise ValueError("Respuesta sin contenido estructurado")
        return msg.parsed, resp.usage
    return _do()

# ──────────────────────────────────────────────────────────────────────────────
# 7. Postproceso: limpiar ítems nulos y el bug "Original"
# ──────────────────────────────────────────────────────────────────────────────


def clean_phrase(s: str) -> str:
    s = re.sub(r"\s+", " ", str(s)).strip(" \t\n.;:,-•*")
    return s[:1].upper() + s[1:] if s else s


def _is_null_label(s: str) -> bool:
    """True si la etiqueta es un placeholder vacío o el bug 'Original'."""
    return s.strip().lower() in _NULL_LABELS


def dedupe(items: list[str]) -> list[str]:
    seen, out = set(), []
    for x in items:
        k = x.lower()
        if x and k not in seen:
            seen.add(k)
            out.append(x)
    return out


def postprocess_items(items: list[str]) -> list[str]:
    """Limpia frases: elimina nulls/placeholders, capitaliza, deduplica."""
    cleaned = [clean_phrase(x) for x in items]
    return dedupe([x for x in cleaned if x and not _is_null_label(x)])


def postprocess_actions(actions: list[dict]) -> list[dict]:
    """Limpia acciones: elimina las que tienen accion='Original' u otro null."""
    out = []
    for a in actions:
        actor = clean_phrase(a.get("actor", "")) or "No especificado"
        accion = clean_phrase(a.get("accion", ""))
        if accion and not _is_null_label(accion):
            out.append({"actor": actor, "accion": accion})
    # Deduplicar por (actor, accion)
    seen, deduped = set(), []
    for a in out:
        key = (a["actor"].lower(), a["accion"].lower())
        if key not in seen:
            seen.add(key)
            deduped.append(a)
    return deduped


def apply_cap(items: list, max_items: Optional[int]) -> tuple[list, int]:
    """
    Recorta una lista a max_items (tope de seguridad de --max-items).
    Regresa la lista (recortada o no) y cuántos ítems traía antes.
    Si max_items es None, no recorta.
    """
    n = len(items)
    if max_items is None or n <= max_items:
        return items, n
    return items[:max_items], n

# ──────────────────────────────────────────────────────────────────────────────
# 8. Extracción por entrevista
# ──────────────────────────────────────────────────────────────────────────────


def analyze_transcript(
    text: str,
    metadata: InterviewMetadata,
    client,
    model: str,
    tracker: UsageTracker,
    prompt_variant: str = "p1_baseline",
    max_items: Optional[int] = None,
    temperature: Optional[float] = 0.0,
    esquema: str = "con_largo",
) -> InterviewResult:
    result = InterviewResult(metadata=metadata, prompt_variant=prompt_variant)
    info = MODEL_CATALOG.get(model)
    system = build_prompt(prompt_variant)
    est_in = count_tokens(text) + count_tokens(system) + PROMPT_OVERHEAD_TOKENS
    if info and est_in > CONTEXT_SAFETY * info.context:
        result.error = (
            f"La entrada (~{est_in:,} tokens) excede el {int(CONTEXT_SAFETY*100)}% "
            f"del context window de {model} ({info.context:,})."
        )
        log.error(f"  {metadata.Id_entrevista}: {result.error}")
        return result
    try:
        parsed, usage = _call_structured(
            client, model, system,
            USER_TEMPLATE.format(transcript=text),
            ESQUEMAS[esquema], temperature,
        )
        result.input_tokens, result.output_tokens = tracker.add(usage)

        # Postproceso: limpia nulls y el bug "Original"
        result.Preocupacion_principal_acerca_del_agua = postprocess_items(
            parsed.Preocupacion_principal_acerca_del_agua
        )
        result.Causas_principales = postprocess_items(
            parsed.Causas_principales)
        result.Consecuencias = postprocess_items(parsed.Consecuencias)
        result.Acciones_y_Actores_que_realizan_dichas_acciones = postprocess_actions(
            [{"actor": a.actor, "accion": a.accion}
             for a in parsed.Acciones_y_Actores_que_realizan_dichas_acciones]
        )

        # Tope de seguridad (--max-items): se aplica después de limpiar y deduplicar
        for campo in ("Preocupacion_principal_acerca_del_agua", "Causas_principales",
                      "Consecuencias", "Acciones_y_Actores_que_realizan_dichas_acciones"):
            recortada, n = apply_cap(getattr(result, campo), max_items)
            if n > len(recortada):
                result.items_antes_de_tope[campo] = n
                log.warning(
                    f"  {metadata.Id_entrevista}: {campo} recortado de {n} a {max_items}")
            setattr(result, campo, recortada)
    except Exception as exc:
        log.error(f"  Error analizando {metadata.Id_entrevista}: {exc}")
        result.error = str(exc)
    return result


# ──────────────────────────────────────────────────────────────────────────────
# 9. Salidas: CSV consolidado
# ──────────────────────────────────────────────────────────────────────────────
COL_PREOC = "Preocupacion_principal_acerca_del_agua"
COL_CAUSAS = "Causas_principales"
COL_CONSEC = "Consecuencias"
COL_ACC = "Acciones_y_Actores_que_realizan_dichas_acciones"


def _fmt_action(a: dict) -> str:
    return f"{a['actor']}: {a['accion']}"


def results_to_dataframe(results: list[InterviewResult]) -> pd.DataFrame:
    rows = []
    for r in sorted(
        (r for r in results if not r.error),
        key=lambda r: r.metadata.Id_entrevista,
    ):
        rows.append({
            "Id_entrevista": r.metadata.Id_entrevista,
            "prompt_variant": r.prompt_variant,
            COL_PREOC:  "\n".join(r.Preocupacion_principal_acerca_del_agua),
            COL_CAUSAS: "\n".join(r.Causas_principales),
            COL_CONSEC: "\n".join(r.Consecuencias),
            COL_ACC:    "\n".join(_fmt_action(a)
                                  for a in r.Acciones_y_Actores_que_realizan_dichas_acciones),
        })
    return pd.DataFrame(rows)


_checkpoint_lock = threading.Lock()


def _csv_name(prompt_variant: str) -> str:
    """Nombre canónico del CSV de salida para una variante dada."""
    return f"analysis_v7_{prompt_variant}.csv"


def make_run_dir(*parts: str) -> Path:
    """
    Arma la carpeta de una corrida nueva dentro de OUTPUTS_ROOT:
    outputs_v<versión>/<fecha_hora>_<partes>/  ej. 20261026_101530_p1_baseline_gpt-4o-mini
    Si ya existe (dos corridas en el mismo segundo), agrega _2, _3...
    No la crea: eso lo hacen run_pipeline y count_only.
    """
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = OUTPUTS_ROOT / "_".join([stamp, *parts])
    path, n = base, 2
    while path.exists():
        path = base.with_name(f"{base.name}_{n}")
        n += 1
    return path


def resolve_resume_dir(path: str) -> Path:
    """
    Valida la carpeta que se pasa con --out (solo para retomar una corrida cortada).
    Acepta la ruta completa o solo el nombre de la carpeta dentro de OUTPUTS_ROOT.
    Debe existir y estar dentro de OUTPUTS_ROOT; si no, el script se detiene.
    """
    root = OUTPUTS_ROOT.resolve()
    p = Path(path)
    if not p.exists() and not p.is_absolute():
        p = OUTPUTS_ROOT / p
    p = p.resolve()
    if not p.is_dir():
        raise SystemExit(f"--out: la carpeta no existe: {p}")
    if p == root or not p.is_relative_to(root):
        raise SystemExit(
            f"--out solo sirve para retomar corridas dentro de {root}\n"
            f"Carpeta recibida: {p}"
        )
    return p


def check_resume_config(config_path: Path, current: dict) -> None:
    """
    Al retomar una corrida, compara su experiment_config.json con la configuración
    actual. Si cambió algo que afecta los resultados, el script se detiene para
    no mezclar condiciones en la misma carpeta.
    """
    saved = json.loads(config_path.read_text(encoding="utf-8"))
    keys = ("script_version", "experimento", "modelo", "temperatura",
            "prompt_variant", "max_items_postproceso", "prompt_text", "esquema")
    diffs = [k for k in keys if saved.get(k) != current.get(k)]
    if diffs:
        raise SystemExit(
            f"No se puede retomar {config_path.parent}: cambió {', '.join(diffs)} "
            f"respecto a la corrida original.\n"
            f"Corre sin --out para empezar una corrida nueva."
        )


def save_checkpoints(result: InterviewResult, all_results: list[InterviewResult], out_dir: Path):
    d = out_dir / "individual"
    d.mkdir(parents=True, exist_ok=True)
    (d / f"{result.metadata.Id_entrevista}.json").write_text(
        json.dumps(asdict(result), ensure_ascii=False, indent=2), encoding="utf-8"
    )
    df = results_to_dataframe(all_results)
    if not df.empty:
        df.to_csv(out_dir / _csv_name(result.prompt_variant),
                  index=False, encoding="utf-8-sig")


def load_existing_result(out_dir: Path, interview_id: str) -> Optional[InterviewResult]:
    p = out_dir / "individual" / f"{interview_id}.json"
    if not p.exists():
        return None
    try:
        d = json.loads(p.read_text(encoding="utf-8"))
        if d.get("error"):
            return None
        d["metadata"] = InterviewMetadata(**d["metadata"])
        return InterviewResult(**d)
    except Exception:
        return None

# ──────────────────────────────────────────────────────────────────────────────
# 10. Pipeline principal
# ──────────────────────────────────────────────────────────────────────────────


def _process_single_file(
    fpath: Path, client, model: str, out_path: Path,
    all_results: list[InterviewResult], tracker: UsageTracker,
    resume: bool, prompt_variant: str, max_items: Optional[int],
    temperature: float = 0.0,
    esquema: str = "con_largo",
):
    metadata = parse_interview_id(fpath.name)
    if resume:
        prev = load_existing_result(out_path, metadata.Id_entrevista)
        if prev:
            log.info(f"Omitida (ya procesada): {fpath.name}")
            with _checkpoint_lock:
                all_results.append(prev)
            return
    log.info(f"Procesando: {fpath.name}")
    raw_text = load_transcript(fpath)
    if not raw_text.strip():
        result = InterviewResult(
            metadata=metadata, prompt_variant=prompt_variant, error="Archivo vacío"
        )
    else:
        result = analyze_transcript(
            preprocess(raw_text), metadata, client, model, tracker,
            prompt_variant=prompt_variant, max_items=max_items,
            temperature=temperature,
            esquema=esquema,
        )
    with _checkpoint_lock:
        all_results.append(result)
        save_checkpoints(result, all_results, out_path)


def make_client():
    """Crea el cliente de OpenAI y verifica la llave antes de crear carpetas o gastar."""
    import openai
    key = os.environ.get("OPENAI_API_KEY")
    if not key:
        raise SystemExit("Falta OPENAI_API_KEY en el archivo .env")
    client = openai.OpenAI(api_key=key)
    try:
        client.models.list()  # llamada mínima: no consume tokens
    except openai.AuthenticationError as exc:
        raise SystemExit(f"La llave de OpenAI no es válida o expiró:\n{exc}")
    except openai.APIConnectionError as exc:
        raise SystemExit(f"No hay conexión con OpenAI:\n{exc}")
    except openai.PermissionDeniedError:
        log.warning(
            "No se pudo verificar la llave (permisos restringidos); se continúa.")
        return client
    log.info("Llave de OpenAI verificada.")
    return client


def run_pipeline(
    input_folder: str,
    output_folder: str,
    model: str,
    prompt_variant: str = "p1_baseline",
    max_items: Optional[int] = None,
    max_workers: int = 2,
    resume: bool = True,
    client=None,
    temperature: float = 0.0,
    experimento: Optional[str] = None,
    esquema: str = "con_largo",
) -> list[InterviewResult]:
    """
    Ejecuta la extracción sobre todas las transcripciones de `input_folder`
    y guarda los resultados en `output_folder/analysis_v7_{prompt_variant}.csv`.
    """
    client = client or make_client()
    out_path = Path(output_folder)
    out_path.mkdir(parents=True, exist_ok=True)
    tracker = UsageTracker(model)

    # La temperatura solo se manda a los modelos que la admiten (TEMPERATURE_PREFIXES)
    temp_aplicada = temperature if model.startswith(
        TEMPERATURE_PREFIXES) else None
    if temp_aplicada is None:
        log.warning(
            f"{model} no está en TEMPERATURE_PREFIXES: no se le manda temperatura.")
    log.info(
        f"Experimento: {experimento} | Modelo: {model} | Prompt variant: {prompt_variant} | "
        f"max_items: {max_items} | temperatura: {temp_aplicada} | esquema: {esquema}")
    log.info(f"Salida: {out_path.resolve()}")

    # Guardar metadatos del experimento
    experiment_meta = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "script_version": SCRIPT_VERSION,
        "experimento": experimento,
        "modelo": model,
        "temperatura": temp_aplicada,
        "prompt_variant": prompt_variant,
        "max_items_postproceso": max_items,
        "esquema": esquema,
        "esquema_json": ESQUEMAS[esquema].model_json_schema(),
        "prompt_text": build_prompt(prompt_variant),
        "prompts_dir": str(_ACTIVE_PROMPTS_DIR.resolve()),
        "carpeta_transcripciones": str(Path(input_folder).resolve()),
    }
    config_path = out_path / "experiment_config.json"
    if resume and config_path.exists():
        # Retomar: la configuración debe ser idéntica a la original (no se reescribe)
        check_resume_config(config_path, experiment_meta)
        log.info(f"Retomando corrida existente: {out_path.resolve()}")
    else:
        config_path.write_text(
            json.dumps(experiment_meta, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    files = list_transcripts(input_folder)
    if not files:
        raise SystemExit(
            f"No se encontraron transcripciones en {input_folder}")
    log.info(f"Transcripciones encontradas: {len(files)}")

    all_results: list[InterviewResult] = []
    with ThreadPoolExecutor(max_workers=max_workers) as executor:
        futures = {
            executor.submit(
                _process_single_file, f, client, model, out_path,
                all_results, tracker, resume, prompt_variant, max_items,
                temperature, esquema,
            ): f
            for f in files
        }
        for fut in tqdm(as_completed(futures), total=len(futures), desc=f"[{prompt_variant}]"):
            try:
                fut.result()
            except Exception as exc:
                log.error(f"Error fatal en {futures[fut].name}: {exc}")

    failed = [r.metadata.Id_entrevista for r in all_results if r.error]
    if failed:
        log.warning(f"{len(failed)} entrevistas con error: {failed}")

    df = results_to_dataframe(all_results)
    if not df.empty:
        out_csv = out_path / _csv_name(prompt_variant)
        df.to_csv(out_csv, index=False, encoding="utf-8-sig")
        log.info(
            f"CSV consolidado: {out_csv.resolve()} ({len(df)} entrevistas)")
    else:
        log.warning("No se generaron resultados válidos; el CSV no fue creado.")

    # Registro acumulativo de uso
    usage = {
        "fecha": datetime.now().isoformat(timespec="seconds"),
        "modelo": model,
        "prompt_variant": prompt_variant,
        "csv_salida": _csv_name(prompt_variant),
        "tokens_entrada": tracker.input,
        "tokens_salida": tracker.output,
        "costo_usd_estimado": tracker.cost(),
        "entrevistas_recortadas": [r.metadata.Id_entrevista for r in all_results if r.items_antes_de_tope],
    }
    with open(out_path / "usage_log_v7.jsonl", "a", encoding="utf-8") as fh:
        fh.write(json.dumps(usage, ensure_ascii=False) + "\n")
    log.info(
        f"Uso de tokens y costo estimado:\n"
        + json.dumps(usage, ensure_ascii=False, indent=2)
    )
    return all_results


def run_all_variants(
    input_folder: str,
    base_out_dir: str,
    model: str,
    max_items: Optional[int] = None,
    max_workers: int = 2,
    resume: bool = True,
    client=None,
    temperature: float = 0.0,
    experimento: Optional[str] = None,
    esquema: str = "con_largo",
) -> dict[str, list[InterviewResult]]:
    """Ejecuta todas las variantes de prompt en secuencia. Cada una va a su propio subdirectorio."""
    # Revisa los .md de todas las variantes antes de gastar en la API
    for v in PROMPT_VARIANTS:
        build_prompt(v)
    client = client or make_client()
    results_by_variant = {}
    for variant in PROMPT_VARIANTS:
        out = str(Path(base_out_dir) / variant)
        log.info(f"\n{'='*60}\nVariante: {variant}\n{'='*60}")
        results_by_variant[variant] = run_pipeline(
            input_folder, out, model,
            prompt_variant=variant, max_items=max_items,
            max_workers=max_workers, resume=resume, client=client,
            temperature=temperature, experimento=experimento, esquema=esquema,
        )
    # Resumen comparativo de costos
    print("\n" + "="*60)
    print("RESUMEN DE COSTO POR VARIANTE")
    print("="*60)
    for variant in PROMPT_VARIANTS:
        log_path = Path(base_out_dir) / variant / "usage_log_v7.jsonl"
        if log_path.exists():
            last = json.loads(log_path.read_text(
                encoding="utf-8").strip().split("\n")[-1])
            print(f"  {variant:25s}  ${last.get('costo_usd_estimado', 0):.4f} USD"
                  f"  →  {last.get('csv_salida', '?')}")
    return results_by_variant


# ──────────────────────────────────────────────────────────────────────────────
# 11. CLI
# ──────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    ap = argparse.ArgumentParser(
        description="Pipeline MEGADAPT v7.3 — framework de experimentación con variantes de prompt",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog=f"""
Variantes de prompt disponibles (un .md por variante en PROMPTS/):
{chr(10).join(f'  {k}' for k in PROMPT_VARIANTS)}

Ejemplos de uso (cada comando va en una sola línea):

  # Estimar costo sin llamar a la API:
  python megadapt_qualitative_pipeline_v7.3.py --folder ./transcripciones --count-only

  # Correr una variante (lee PROMPTS/p1_baseline.md):
  python megadapt_qualitative_pipeline_v7.3.py --folder ./transcripciones --prompt-variant p1_baseline
  # → salida: outputs_v{SCRIPT_VERSION}/<fecha_hora>_p1_baseline_gpt-4o-mini/analysis_v7_p1_baseline.csv

  # Con tope de seguridad de 50 ítems por campo (se aplica en el postproceso):
  python megadapt_qualitative_pipeline_v7.3.py --folder ./transcripciones --prompt-variant p1_baseline --max-items 50

  # Con etiqueta de experimento y temperatura:
  python megadapt_qualitative_pipeline_v7.3.py --folder ./transcripciones --prompt-variant p1_baseline --experimento E1 --temperatura 0.0
  # → salida: outputs_v{SCRIPT_VERSION}/<fecha_hora>_E1_p1_baseline_gpt-4o-mini/

  # Usar una carpeta de prompts distinta a PROMPTS/:
  python megadapt_qualitative_pipeline_v7.3.py --folder ./transcripciones --prompt-variant p2_definiciones --prompts-dir ./otra_carpeta

  # Correr las tres variantes en secuencia (una subcarpeta por variante):
  python megadapt_qualitative_pipeline_v7.3.py --folder ./transcripciones --run-all-variants

  # Retomar una corrida que se cortó (basta el nombre de su carpeta):
  python megadapt_qualitative_pipeline_v7.3.py --folder ./transcripciones --prompt-variant p1_baseline --out 20261026_101530_p1_baseline_gpt-4o-mini

  # Con el formulario sin largo fijo:
  python megadapt_qualitative_pipeline_v7.3.py --prompt-variant p1_baseline --experimento E7 --esquema sin_largo

  # Evaluar todas las corridas (script de evaluación v4):
  python megadapt_eval_nlp_v4.py --runs outputs_v{SCRIPT_VERSION} --metrics bertscore chrf
""",
    )
    ap.add_argument(
        "--folder", default=os.environ.get("MEGADAPT_TRANSCRIPCIONES"),
        help="Carpeta con transcripciones (.docx/.txt/.md). Si no se pone, se usa "
             "MEGADAPT_TRANSCRIPCIONES del archivo .env.")
    ap.add_argument("--out", default=None,
                    help=(f"Solo para retomar una corrida cortada: su carpeta dentro de "
                          f"outputs_v{SCRIPT_VERSION}/ (basta el nombre). Sin --out, cada "
                          f"corrida crea su propia carpeta nueva ahí."))
    ap.add_argument("--model", default="gpt-4o-mini",
                    help=f"Modelo de extracción. Opciones: {', '.join(MODEL_CATALOG)}")
    ap.add_argument("--prompt-variant", default="p1_baseline",
                    choices=PROMPT_VARIANTS,
                    help="Variante de prompt a usar (default: p1_baseline)")
    ap.add_argument("--prompts-dir", default=None,
                    help=(
                        "Carpeta con un .md por variante. "
                        "Por defecto PROMPTS/ junto al script; si no existe, el script se detiene."
                    ))
    ap.add_argument("--max-items", type=int, default=None,
                    help="Tope de ítems por campo, aplicado en el postproceso (no va en el prompt). Sin valor: no recorta.")
    ap.add_argument("--experimento", default=None,
                    help="Etiqueta del experimento (ej. E1). Queda en experiment_config.json "
                         "y en el nombre de la carpeta, después de la fecha. Solo letras, números, - y _.")
    ap.add_argument("--temperatura", type=float, default=0.0,
                    help="Temperatura del modelo (default: 0.0). Solo se aplica a gpt-4o y gpt-4.1.")
    ap.add_argument("--esquema", default="con_largo", choices=list(ESQUEMAS),
                    help="Formulario de salida: con_largo (descripciones con '4 o 5 palabras', el de E1) "
                         "o sin_largo (sin número de palabras). Default: con_largo.")
    ap.add_argument("--workers", type=int, default=2,
                    help="Hilos concurrentes (ajustar según rate limit, default: 2)")
    ap.add_argument("--count-only", action="store_true",
                    help="Solo cuenta tokens/costo; no llama a la API")
    ap.add_argument("--models", default=",".join(MODEL_CATALOG),
                    help="Modelos a comparar en --count-only (separados por coma)")
    ap.add_argument("--no-resume", action="store_true",
                    help="Reprocesar todas las entrevistas aunque ya existan sus JSONs")
    ap.add_argument("--run-all-variants", action="store_true",
                    help="Ejecutar todas las variantes en secuencia; cada una queda en una subcarpeta de la corrida.")
    args = ap.parse_args()

    # Configurar directorio de prompts (antes de cualquier llamada a build_prompt)
    set_prompts_dir(args.prompts_dir)

    # La etiqueta del experimento va en el nombre de la carpeta: solo caracteres seguros
    if args.experimento and not re.fullmatch(r"[A-Za-z0-9_-]+", args.experimento):
        ap.error("--experimento solo admite letras, números, - y _ (ej. E1, E1_rep2)")
    prefijo = [args.experimento] if args.experimento else []

    if args.count_only:
        if not args.folder:
            ap.error("--count-only requiere --folder")
        if args.out:
            ap.error(
                "--out solo sirve para retomar corridas; no aplica con --count-only")
        unknown = [m for m in args.models.split(",") if m not in MODEL_CATALOG]
        if unknown:
            ap.error(f"Modelos fuera del catálogo: {unknown}")
        count_only(args.folder, str(make_run_dir("conteo", args.prompt_variant)),
                   args.models.split(","), variant=args.prompt_variant)
    elif args.run_all_variants:
        if not args.folder:
            ap.error("--run-all-variants requiere --folder")
        run_all_variants(
            args.folder,
            str(resolve_resume_dir(args.out)) if args.out else str(
                make_run_dir(*prefijo, "todas", args.model)),
            args.model,
            max_items=args.max_items,
            max_workers=args.workers,
            resume=not args.no_resume,
            temperature=args.temperatura,
            experimento=args.experimento,
            esquema=args.esquema,
        )
    else:
        if not args.folder:
            ap.error("Falta --folder con la ruta a las transcripciones")
        out_dir = (str(resolve_resume_dir(args.out)) if args.out
                   else str(make_run_dir(*prefijo, args.prompt_variant, args.model)))
        run_pipeline(
            args.folder, out_dir, args.model,
            prompt_variant=args.prompt_variant,
            max_items=args.max_items,
            max_workers=args.workers,
            resume=not args.no_resume,
            temperature=args.temperatura,
            experimento=args.experimento,
            esquema=args.esquema,
        )
