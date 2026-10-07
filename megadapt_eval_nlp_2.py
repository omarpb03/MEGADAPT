"""
MEGADAPT - Evaluación NLP de la codificación cualitativa (v3)

Calcula BLEU, chrF++ y BERTScore para comparar la salida del pipeline
contra la codificación humana (hoja «Integracion» del Excel).

Estas métricas complementan las métricas léxicas/LLM de megadapt_compare_human.py:
  • BLEU:    precisión de n-gramas de palabras (sacrebleu).
  • chrF++:  F-score de n-gramas de caracteres Y palabras (sacrebleu).
  • BERTScore: similitud semántica mediante embeddings contextuales.

─── Por qué adaptamos las métricas ───────────────────────────────────────────
BLEU y chrF++ fueron diseñadas para traducción (una hipótesis vs. una referencia).
Aquí tenemos CONJUNTOS de frases cortas por campo y entrevista. La adaptación:

  1. Recall (cobertura del humano):
     Para cada ítem humano, buscamos el ítem del modelo con mayor score
     y tomamos ese máximo. El recall final es la media de esos máximos.
     → Mide: ¿cuánto de lo que dijo el humano captura el modelo?

  2. Precision (respaldo del modelo):
     Para cada ítem del modelo, buscamos el ítem humano con mayor score.
     → Mide: ¿cuánto de lo que dijo el modelo tiene respaldo humano?

  3. F1: media armónica de recall y precision.

─── Modos de uso ─────────────────────────────────────────────────────────────

# Resumen por corrida (genera CSVs + Excel directivo):
python megadapt_eval_nlp.py \
    --human "2.Análisis Cualitativo Integración.xlsx" \
    --llm v5=output_v5/analysis_standardized_v5.csv \
        v6=output_v6/analysis_standardized_v6.csv \
    --out ./eval_nlp --metrics bleu chrf --xlsx

# Enriquecer el CSV del pipeline:
python megadapt_eval_nlp.py \
    --human "2.Análisis Cualitativo Integración.xlsx" \
    --llm output_v5/analysis_standardized_v5.csv \
    --out ./eval_nlp --metrics bleu chrf --enrich

# Con BERTScore (descarga ~440 MB la primera vez):
python megadapt_eval_nlp.py \
    --human "2.Análisis Cualitativo Integración.xlsx" \
    --llm v5=output_v5/analysis_standardized_v5.csv \
    --out ./eval_nlp --metrics bleu chrf bertscore \
    --bertscore-model dccuchile/bert-base-spanish-wwm-cased \
    --lang es --xlsx

─── Dependencias ─────────────────────────────────────────────────────────────
pip install sacrebleu bert-score pandas openpyxl
(torch se instala como dependencia de bert-score)

─── Interpretación ───────────────────────────────────────────────────────────
• BLEU varía 0–1. Valores < 0.20 son normales para frases cortas y resúmenes.
• chrF++ varía 0–1. Más robusto que BLEU para español (morfología rica).
• BERTScore F1 varía ~0.7–1.0. > 0.85 indica alta similitud semántica.

─── Excel directivo (--xlsx) ────────────────────────────────────────────────
Genera megadapt_evaluacion.xlsx con dos hojas:

  Hoja 1 "Benchmarking Modelos"
      Pivot corrida × categoría mostrando el F1 de la métrica primaria
      (bertscore > chrf > bleu, según lo que se haya calculado).
      Mapa de calor: Rojo(<0.40) · Amarillo(~0.625) · Verde(>0.85).

  Hoja 2 "Matriz Cruda"
      Datos por entrevista. Columnas _recall y _precision ocultas.
      Columnas _f1 con el mismo mapa de calor.
"""
from __future__ import annotations

import argparse
import logging
import re
import unicodedata
from collections import defaultdict
from pathlib import Path
from typing import Optional

import pandas as pd

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s  %(levelname)-8s  %(message)s",
    datefmt="%H:%M:%S",
)
log = logging.getLogger(__name__)

# ──────────────────────────────────────────────────────────────────────────────
# Constantes
# ──────────────────────────────────────────────────────────────────────────────
COLUMNS = ["preocupacion", "causa", "consecuencia", "accion"]
LLM_CSV_COLS = {
    "preocupacion": "Preocupacion_principal_acerca_del_agua",
    "causa": "Causas_principales",
    "consecuencia": "Consecuencias",
    "accion": "Acciones_y_Actores_que_realizan_dichas_acciones",
}
HUMAN_HEADER_KEYS = {
    "preocupacion": "preocupaci",
    "causa": "causas principales",
    "consecuencia": "consecuencias",
    "accion": "acciones",
}
NULL_ITEMS = {"ninguno", "ninguna", "none", "na", "n a", "no aplica", "nan", "-"}

# Colores para el mapa de calor del Excel (escala Rojo→Amarillo→Verde)
_HEATMAP_LOW   = "F8696B"   # rojo   → F1 ≤ 0.40
_HEATMAP_MID   = "FFEB84"   # amarillo → F1 ≈ 0.625
_HEATMAP_HIGH  = "63BE7B"   # verde  → F1 ≥ 0.85
_HEATMAP_THRESH = (0.40, 0.625, 0.85)

# ──────────────────────────────────────────────────────────────────────────────
# Normalización
# ──────────────────────────────────────────────────────────────────────────────
def strip_accents(s: str) -> str:
    return "".join(
        c for c in unicodedata.normalize("NFD", str(s))
        if unicodedata.category(c) != "Mn"
    )


def split_items(text) -> list[str]:
    if text is None or (isinstance(text, float) and pd.isna(text)):
        return []
    parts = re.split(r"[\n;/]+|(?:^|\s)\d\.\s+", str(text))
    out = []
    for p in parts:
        p = re.sub(r"\s+", " ", p).strip(" \t.,:-•*")
        if len(p) >= 3 and strip_accents(p).lower().strip() not in NULL_ITEMS:
            out.append(p)
    return out


def code_key(code: str, drop_suffix: bool = False) -> str:
    c = strip_accents(str(code)).upper()
    c = re.sub(r"[_\s]*\d+\s*(RA|DA|TA|ER)?\s*PARTE", "", c)
    c = re.sub(r"[^A-Z0-9]", "", c)
    return re.sub(r"[A-Z]+$", "", c) if drop_suffix else c


# ──────────────────────────────────────────────────────────────────────────────
# Carga de datos
# ──────────────────────────────────────────────────────────────────────────────
def load_human(path: str, sheet: str = "Integracion") -> dict[str, dict[str, list[str]]]:
    raw = pd.read_excel(path, sheet_name=sheet, header=None, dtype=str)
    hdr = None
    for i in range(min(10, len(raw))):
        if any(
            isinstance(v, str) and strip_accents(v).lower().strip().startswith("codigo")
            for v in raw.iloc[i]
        ):
            hdr = i
            break
    if hdr is None:
        raise ValueError(
            f"No encontré la fila de encabezados (columna 'Código') en la hoja {sheet}"
        )
    header = [
        strip_accents(str(v)).lower() if isinstance(v, str) else ""
        for v in raw.iloc[hdr]
    ]
    code_idx = next(i for i, h in enumerate(header) if h.strip().startswith("codigo"))
    idx = {}
    for col, key in HUMAN_HEADER_KEYS.items():
        idx[col] = next((i for i, h in enumerate(header) if key in h), None)
        if idx[col] is None:
            raise ValueError(f"No encontré columna '{key}' en: {header}")
    extra_acc = list(range(idx["accion"] + 1, raw.shape[1]))
    data: dict[str, dict[str, list[str]]] = {}
    current = None
    for i in range(hdr + 1, len(raw)):
        row = raw.iloc[i]
        code = row.iloc[code_idx]
        if isinstance(code, str) and code.strip():
            current = code.strip()
            data.setdefault(current, {c: [] for c in COLUMNS})
        if current is None:
            continue
        for col in COLUMNS:
            cells = [row.iloc[idx[col]]] + (
                [row.iloc[j] for j in extra_acc] if col == "accion" else []
            )
            for cell in cells:
                data[current][col] += split_items(cell)
    return data


def load_llm(path: str) -> dict[str, dict[str, list[str]]]:
    df = pd.read_csv(path, dtype=str, encoding="utf-8-sig").fillna("")
    out: dict[str, dict[str, list[str]]] = {}
    for _, r in df.iterrows():
        base = str(r["Id_entrevista"])
        d = out.setdefault(base, {c: [] for c in COLUMNS})
        for col in COLUMNS:
            d[col] += [
                x.strip()
                for x in str(r[LLM_CSV_COLS[col]]).split("\n")
                if x.strip()
            ]
    return out


def match_ids(
    human: dict, llm: dict
) -> tuple[dict[str, list[str]], list[str], list[str]]:
    by_key, by_key_short = defaultdict(list), defaultdict(list)
    for lid in llm:
        by_key[code_key(lid)].append(lid)
        by_key_short[code_key(lid, True)].append(lid)
    pairs, used = {}, set()
    for hid in human:
        cand = by_key.get(code_key(hid)) or by_key_short.get(code_key(hid, True))
        if cand:
            pairs[hid] = cand
            used.update(cand)
    no_h = [h for h in human if h not in pairs]
    no_l = [l for l in llm if l not in used]
    return pairs, no_h, no_l


def reverse_pairs(pairs: dict[str, list[str]]) -> dict[str, str]:
    rev = {}
    for hid, lids in pairs.items():
        for lid in lids:
            rev[lid] = hid
    return rev


# ──────────────────────────────────────────────────────────────────────────────
# MÉTRICAS NLP
# ──────────────────────────────────────────────────────────────────────────────

def _sacrebleu_sentence_bleu(hyp: str, ref: str) -> float:
    from sacrebleu.metrics import BLEU
    bleu = BLEU(smooth_method="exp", effective_order=True)
    return bleu.sentence_score(hyp, [ref]).score


def _sacrebleu_sentence_chrf(hyp: str, ref: str, word_order: int = 2) -> float:
    from sacrebleu.metrics import CHRF
    chrf = CHRF(word_order=word_order)
    return chrf.sentence_score(hyp, [ref]).score


def _pairwise_max_f1(
    human_items: list[str],
    model_items: list[str],
    score_fn,
) -> dict:
    if not human_items or not model_items:
        return {"recall": None, "precision": None, "f1": None}

    matrix = [
        [score_fn(m, h) for m in model_items]
        for h in human_items
    ]
    recall = sum(max(row) for row in matrix) / len(human_items) / 100.0
    precision = (
        sum(max(matrix[i][j] for i in range(len(human_items)))
            for j in range(len(model_items)))
        / len(model_items) / 100.0
    )
    f1 = (
        2 * recall * precision / (recall + precision)
        if (recall + precision) > 0
        else 0.0
    )
    return {"recall": round(recall, 4), "precision": round(precision, 4), "f1": round(f1, 4)}


def compute_bleu(human_items: list[str], model_items: list[str]) -> dict:
    return _pairwise_max_f1(human_items, model_items, _sacrebleu_sentence_bleu)


def compute_chrf(human_items: list[str], model_items: list[str]) -> dict:
    return _pairwise_max_f1(human_items, model_items, _sacrebleu_sentence_chrf)


class _BertScoreCache:
    _instance: Optional["_BertScoreCache"] = None

    def __init__(self, model_type: str, lang: str, device: str):
        self.model_type = model_type
        self.lang = lang
        self.device = device

    @classmethod
    def get(cls, model_type: str, lang: str, device: str) -> "_BertScoreCache":
        if cls._instance is None:
            cls._instance = cls(model_type, lang, device)
        return cls._instance

    def score_pairwise(
        self, cands: list[str], refs: list[str]
    ) -> tuple[list[float], list[float], list[float]]:
        from bert_score import score as bs_score
        P, R, F = bs_score(
            cands, refs,
            model_type=self.model_type,
            lang=self.lang,
            device=self.device,
            verbose=False,
        )
        return P.tolist(), R.tolist(), F.tolist()


def compute_bertscore(
    human_items: list[str],
    model_items: list[str],
    cache: "_BertScoreCache",
) -> dict:
    if not human_items or not model_items:
        return {"recall": None, "precision": None, "f1": None}

    pairs_cand, pairs_ref = [], []
    for h in human_items:
        for m in model_items:
            pairs_cand.append(m)
            pairs_ref.append(h)

    _, _, F = cache.score_pairwise(pairs_cand, pairs_ref)
    n_h, n_m = len(human_items), len(model_items)
    matrix = [F[i * n_m:(i + 1) * n_m] for i in range(n_h)]

    recall = sum(max(row) for row in matrix) / n_h
    precision = (
        sum(max(matrix[i][j] for i in range(n_h)) for j in range(n_m)) / n_m
    )
    f1 = (
        2 * recall * precision / (recall + precision)
        if (recall + precision) > 0
        else 0.0
    )
    return {"recall": round(recall, 4), "precision": round(precision, 4), "f1": round(f1, 4)}


def _compute_all(h_items, m_items, metrics, bs_cache):
    result = {}
    if not h_items or not m_items:
        for met in metrics:
            result[met] = {"recall": None, "precision": None, "f1": None}
        return result
    if "bleu" in metrics:
        result["bleu"] = compute_bleu(h_items, m_items)
    if "chrf" in metrics:
        result["chrf"] = compute_chrf(h_items, m_items)
    if "bertscore" in metrics:
        result["bertscore"] = compute_bertscore(h_items, m_items, bs_cache)
    return result


# ──────────────────────────────────────────────────────────────────────────────
# EXCEL DIRECTIVO
# ──────────────────────────────────────────────────────────────────────────────

def _xl_apply_heatmap(ws, range_str: str) -> None:
    """
    Aplica ColorScaleRule de tres paradas al rango indicado.
    Rojo (≤ 0.40) → Amarillo (≈ 0.625) → Verde (≥ 0.85).
    """
    from openpyxl.formatting.rule import ColorScaleRule
    low, mid, high = _HEATMAP_THRESH
    ws.conditional_formatting.add(
        range_str,
        ColorScaleRule(
            start_type="num", start_value=low,  start_color=_HEATMAP_LOW,
            mid_type="num",   mid_value=mid,    mid_color=_HEATMAP_MID,
            end_type="num",   end_value=high,   end_color=_HEATMAP_HIGH,
        ),
    )


def _xl_header_row(ws, row: int, n_cols: int, color: str = "1F3864") -> None:
    """Aplica estilo de encabezado (fondo oscuro, texto blanco, centrado)."""
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    fill = PatternFill("solid", fgColor=color)
    font = Font(color="FFFFFF", bold=True, size=11)
    thin = Side(style="thin", color="A0A0A0")
    brd  = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c in range(1, n_cols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = brd
    ws.row_dimensions[row].height = 28


def export_to_excel(
    df_por_entrevista: pd.DataFrame,
    out_dir: Path,
    metrics: list[str],
    primary_metric: Optional[str] = None,
    filename: str = "megadapt_evaluacion.xlsx",
) -> Optional[Path]:
    """
    Genera un Excel directivo a partir del DataFrame de evaluate().

    Parámetros
    ----------
    df_por_entrevista : DataFrame
        Salida de evaluate() con columnas: corrida, Id_humano, columna,
        n_items_humano, n_items_modelo, {metric}_recall, {metric}_precision,
        {metric}_f1 para cada métrica calculada.
    out_dir : Path
        Directorio de salida.
    metrics : list[str]
        Métricas calculadas (ej. ['bleu', 'chrf', 'bertscore']).
    primary_metric : str, opcional
        Métrica que se muestra en la Hoja 1. Se elige automáticamente
        en orden de prioridad: bertscore > chrf > bleu.
    filename : str
        Nombre del archivo de salida.

    Hojas generadas
    ---------------
    "Benchmarking Modelos"
        Pivot: corrida × categoría → F1 promedio de la métrica primaria.
        Columna TOTAL = media de las 4 categorías.
        Mapa de calor por umbral.

    "Matriz Cruda"
        Todos los datos por entrevista.
        Columnas _recall y _precision ocultas por defecto.
        Columnas _f1 con mapa de calor.

    Devuelve la ruta al archivo generado, o None si openpyxl no está instalado.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.utils.dataframe import dataframe_to_rows
    except ImportError:
        log.warning(
            "openpyxl no encontrado — omitiendo exportación Excel. "
            "Instalar con: pip install openpyxl"
        )
        return None

    # ── Seleccionar métrica primaria ──────────────────────────────────────────
    if primary_metric is None:
        for m in ("bertscore", "chrf", "bleu"):
            if m in metrics:
                primary_metric = m
                break
    if primary_metric not in metrics:
        log.warning(
            f"Métrica primaria '{primary_metric}' no fue calculada. "
            f"Disponibles: {metrics}. Usando la primera disponible."
        )
        primary_metric = metrics[0]

    f1_col = f"{primary_metric}_f1"
    if f1_col not in df_por_entrevista.columns:
        log.error(f"Columna '{f1_col}' no encontrada en el DataFrame. Abortando Excel.")
        return None

    wb = Workbook()

    # ═══════════════════════════════════════════════════════════════════════════
    # HOJA 1 — Benchmarking Modelos
    # ═══════════════════════════════════════════════════════════════════════════
    ws1 = wb.active
    ws1.title = "Benchmarking Modelos"

    # ── Construcción del pivot ────────────────────────────────────────────────
    pivot = (
        df_por_entrevista
        .pivot_table(
            index="corrida",
            columns="columna",
            values=f1_col,
            aggfunc="mean",
        )
        .round(4)
    )
    # Reordenar columnas en el orden canónico del proyecto
    col_order = [c for c in COLUMNS if c in pivot.columns]
    pivot = pivot.reindex(columns=col_order)
    pivot.insert(len(pivot.columns), "TOTAL", pivot.mean(axis=1).round(4))
    pivot = pivot.reset_index()

    # Etiquetas legibles para la presentación
    col_labels = {
        "corrida":       "Corrida / Experimento",
        "preocupacion":  "Preocupación",
        "causa":         "Causa",
        "consecuencia":  "Consecuencia",
        "accion":        "Acción",
        "TOTAL":         "⌀ TOTAL",
    }

    TITLE_ROW = 1
    BLANK_ROW = 2
    HDR_ROW   = 3
    DATA_ROW0 = 4

    # Título
    title_text = (
        f"MEGADAPT — Benchmarking de Corridas   |   "
        f"Métrica: {primary_metric.upper()} F1   |   "
        f"Umbral calidad: 🟢 > {_HEATMAP_THRESH[2]}  "
        f"🟡 {_HEATMAP_THRESH[0]}–{_HEATMAP_THRESH[2]}  "
        f"🔴 < {_HEATMAP_THRESH[0]}"
    )
    ws1.cell(TITLE_ROW, 1).value = title_text
    ws1.cell(TITLE_ROW, 1).font = Font(bold=True, size=12, color="1F3864")
    ws1.row_dimensions[TITLE_ROW].height = 22

    # Encabezados de columna
    for c_idx, col_name in enumerate(pivot.columns, 1):
        ws1.cell(HDR_ROW, c_idx).value = col_labels.get(col_name, col_name)
    _xl_header_row(ws1, HDR_ROW, len(pivot.columns))

    # Datos
    thin_side = Side(style="thin", color="D0D0D0")
    data_border = Border(
        left=thin_side, right=thin_side,
        top=thin_side, bottom=thin_side,
    )
    n_data_rows = len(pivot)
    for r_offset, row_vals in enumerate(pivot.itertuples(index=False)):
        r = DATA_ROW0 + r_offset
        # Fila alternada (cebra ligera)
        row_fill = PatternFill("solid", fgColor="EEF2F8") if r_offset % 2 == 0 else None
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws1.cell(r, c_idx)
            cell.value = val
            cell.border = data_border
            if row_fill:
                cell.fill = row_fill
            if c_idx == 1:
                # Columna "Corrida": negrita, alineado a la izquierda
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="left", vertical="center")
            elif col_labels.get(list(pivot.columns)[c_idx - 1], "") == "⌀ TOTAL":
                # Columna TOTAL: resaltar con fondo azul claro
                cell.fill = PatternFill("solid", fgColor="D6E4F0")
                cell.font = Font(bold=True)
                cell.number_format = "0.000"
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.number_format = "0.000"
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Mapa de calor (excluye columna "corrida" y columna "TOTAL")
    data_end_row = DATA_ROW0 + n_data_rows - 1
    n_cols_total = len(pivot.columns)
    # Categorías: columnas 2 a n_cols_total-1 (excluye TOTAL para que el color sea puro)
    if n_cols_total > 2:
        cat_end_letter = get_column_letter(n_cols_total - 1)
        _xl_apply_heatmap(ws1, f"B{DATA_ROW0}:{cat_end_letter}{data_end_row}")
    # Columna TOTAL con el mismo heatmap pero separada
    total_letter = get_column_letter(n_cols_total)
    _xl_apply_heatmap(ws1, f"{total_letter}{DATA_ROW0}:{total_letter}{data_end_row}")

    # Anchos de columna
    ws1.column_dimensions["A"].width = 34
    for c_idx in range(2, n_cols_total + 1):
        ltr = get_column_letter(c_idx)
        ws1.column_dimensions[ltr].width = 15 if ltr != total_letter else 13

    # Inmovilizar: primera columna + fila de encabezado
    ws1.freeze_panes = f"B{DATA_ROW0}"

    # ═══════════════════════════════════════════════════════════════════════════
    # HOJA 2 — Matriz Cruda
    # ═══════════════════════════════════════════════════════════════════════════
    ws2 = wb.create_sheet("Matriz Cruda")
    all_cols = list(df_por_entrevista.columns)
    n_total_cols = len(all_cols)

    # Escribir header + datos
    for r_idx, row in enumerate(dataframe_to_rows(df_por_entrevista, index=False, header=True), 1):
        ws2.append(row)

    # Estilo de encabezado
    _xl_header_row(ws2, 1, n_total_cols, color="2E4057")

    data_end_r2 = 1 + len(df_por_entrevista)

    # Clasificar columnas: identificar recall/precision (ocultar) y f1 (heatmap)
    hidden_letters: list[str] = []
    f1_col_letters: list[str] = []
    thin_r2 = Side(style="thin", color="E0E0E0")
    brd_r2  = Border(left=thin_r2, right=thin_r2, top=thin_r2, bottom=thin_r2)

    for c_idx, col_name in enumerate(all_cols, 1):
        lname = col_name.lower()
        letter = get_column_letter(c_idx)

        if "_recall" in lname or "_prec" in lname or "_precision" in lname:
            # Ocultar columnas de recall y precision
            hidden_letters.append(letter)
            ws2.column_dimensions[letter].hidden = True

        elif "_f1" in lname:
            f1_col_letters.append(letter)
            # Formato numérico + centrado
            for r in range(2, data_end_r2 + 1):
                cell = ws2.cell(r, c_idx)
                cell.number_format = "0.0000"
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = brd_r2
            ws2.column_dimensions[letter].width = 13

        elif col_name in ("corrida", "Id_humano", "columna"):
            # Columnas de identificación: ancho fijo
            ws2.column_dimensions[letter].width = {
                "corrida": 24, "Id_humano": 18, "columna": 14
            }.get(col_name, 16)
            for r in range(2, data_end_r2 + 1):
                ws2.cell(r, c_idx).border = brd_r2

        else:
            # Columnas numéricas de conteo
            ws2.column_dimensions[letter].width = 12
            for r in range(2, data_end_r2 + 1):
                cell = ws2.cell(r, c_idx)
                cell.alignment = Alignment(horizontal="center")
                cell.border = brd_r2

    # Mapa de calor en columnas F1
    for letter in f1_col_letters:
        _xl_apply_heatmap(ws2, f"{letter}2:{letter}{data_end_r2}")

    # Leyenda de columnas ocultas
    legend_row = data_end_r2 + 2
    ws2.cell(legend_row, 1).value = (
        f"ℹ️  Columnas ocultas ({len(hidden_letters)}): "
        f"_recall y _precision de cada métrica. "
        f"Mostrar desde la pestaña 'Formato > Columnas > Mostrar'."
    )
    ws2.cell(legend_row, 1).font = Font(italic=True, size=9, color="808080")

    # Inmovilizar: Id + columna
    ws2.freeze_panes = "C2"

    # Guardar
    out_path = out_dir / filename
    wb.save(str(out_path))
    log.info(f"Excel directivo guardado en: {out_path.resolve()}")
    return out_path


# ──────────────────────────────────────────────────────────────────────────────
# Modo 1: Resumen por corrida
# ──────────────────────────────────────────────────────────────────────────────
def evaluate(
    human_path: str,
    runs: dict[str, str],
    out_dir: str,
    metrics: list[str],
    sheet: str = "Integracion",
    bertscore_model: str = "microsoft/mdeberta-v3-base",
    lang: str = "es",
    device: str = "cpu",
    xlsx: bool = False,
    primary_metric: Optional[str] = None,
) -> pd.DataFrame:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    human = load_human(human_path, sheet)
    log.info(f"Codificación humana: {len(human)} entrevistas")

    bs_cache = None
    if "bertscore" in metrics:
        log.info(f"Cargando BERTScore (modelo: {bertscore_model}, device: {device})")
        bs_cache = _BertScoreCache.get(bertscore_model, lang, device)

    rows = []
    for run, path in runs.items():
        llm = load_llm(path)
        pairs, no_h, no_l = match_ids(human, llm)
        log.info(
            f"[{run}] emparejadas {len(pairs)} / {len(human)} entrevistas "
            f"({len(no_h)} sin par en el modelo, {len(no_l)} del modelo sin par en humano)"
        )

        for hid, lids in pairs.items():
            for col in COLUMNS:
                h_items = human[hid][col]
                m_items = list(dict.fromkeys(x for lid in lids for x in llm[lid][col]))

                if not h_items:
                    continue

                row: dict = {
                    "corrida": run, "Id_humano": hid, "columna": col,
                    "n_items_humano": len(h_items), "n_items_modelo": len(m_items),
                }

                scores = _compute_all(h_items, m_items, metrics, bs_cache)
                for met, vals in scores.items():
                    row[f"{met}_recall"]    = vals["recall"]
                    row[f"{met}_precision"] = vals["precision"]
                    row[f"{met}_f1"]        = vals["f1"]

                rows.append(row)

    if not rows:
        log.warning("No se generaron filas.")
        return pd.DataFrame()

    df = pd.DataFrame(rows)
    df.to_csv(out / "eval_nlp_por_entrevista.csv", index=False, encoding="utf-8-sig")

    metric_cols = [c for c in df.columns if any(c.startswith(m) for m in metrics)]
    agg = (
        df.groupby(["corrida", "columna"])[metric_cols].mean().round(4)
        .join(df.groupby(["corrida", "columna"]).size().rename("n_entrevistas"))
        .reset_index()
    )
    total = (
        df.groupby("corrida")[metric_cols].mean().round(4)
        .join(df.groupby("corrida")["Id_humano"].nunique().rename("n_entrevistas"))
        .reset_index().assign(columna="TODAS")
    )
    summary = pd.concat([agg, total], ignore_index=True)
    summary.to_csv(out / "eval_nlp_resumen.csv", index=False, encoding="utf-8-sig")

    print("\n" + "=" * 70)
    print("RESUMEN DE MÉTRICAS NLP")
    print("=" * 70)
    print(summary.to_string(index=False))
    print("\nInterpretación rápida:")
    print("  _recall    = cobertura del humano por el modelo  (0–1)")
    print("  _precision = respaldo del modelo en el humano    (0–1)")
    print("  _f1        = media armónica (0–1)")
    print(f"\nArchivos CSV generados en: {out.resolve()}")

    # ── Excel directivo (opcional) ────────────────────────────────────────────
    if xlsx:
        xl_path = export_to_excel(
            df, out, metrics,
            primary_metric=primary_metric,
            filename="megadapt_evaluacion.xlsx",
        )
        if xl_path:
            print(f"\nExcel directivo: {xl_path.resolve()}")

    return summary


# ──────────────────────────────────────────────────────────────────────────────
# Modo 2: Enriquecer CSV del pipeline con columnas de métricas por campo
# ──────────────────────────────────────────────────────────────────────────────
def enrich_csv(
    human_path: str,
    llm_csv: str,
    out_dir: str,
    metrics: list[str],
    sheet: str = "Integracion",
    bertscore_model: str = "microsoft/mdeberta-v3-base",
    lang: str = "es",
    device: str = "cpu",
) -> pd.DataFrame:
    """
    Carga el CSV del pipeline, añade columnas de métricas NLP por campo y
    guarda el resultado como <nombre_original>_con_metricas.csv en out_dir.

    Columnas añadidas por cada métrica M y campo C:
      M_recall_C  M_prec_C  M_f1_C

    Columnas resumen (promedio de los cuatro campos con referencia humana):
      M_recall_TOTAL  M_prec_TOTAL  M_f1_TOTAL

    Filas sin par en la codificación humana → NaN (sin referencia).
    """
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)

    human = load_human(human_path, sheet)
    log.info(f"Codificación humana: {len(human)} entrevistas")

    bs_cache = None
    if "bertscore" in metrics:
        log.info(f"Cargando BERTScore (modelo: {bertscore_model}, device: {device})")
        bs_cache = _BertScoreCache.get(bertscore_model, lang, device)

    df_orig = pd.read_csv(llm_csv, dtype=str, encoding="utf-8-sig").fillna("")
    llm = load_llm(llm_csv)
    pairs, no_h, no_l = match_ids(human, llm)
    rev = reverse_pairs(pairs)

    log.info(
        f"Emparejadas {len(pairs)} / {len(human)} entrevistas del humano "
        f"({len(no_l)} filas del modelo sin par humano → quedarán como NaN)"
    )

    col_names = []
    for met in metrics:
        for campo in COLUMNS:
            col_names += [f"{met}_recall_{campo}", f"{met}_prec_{campo}", f"{met}_f1_{campo}"]
    for met in metrics:
        col_names += [f"{met}_recall_TOTAL", f"{met}_prec_TOTAL", f"{met}_f1_TOTAL"]

    new_cols: dict[str, list] = {c: [] for c in col_names}

    for _, row in df_orig.iterrows():
        lid = str(row["Id_entrevista"])
        hid = rev.get(lid)

        if hid is None:
            for c in col_names:
                new_cols[c].append(None)
            continue

        field_vals: dict[str, dict[str, dict]] = {}
        for col in COLUMNS:
            h_items = human[hid][col]
            m_items = split_items(row.get(LLM_CSV_COLS[col], ""))
            scores = _compute_all(h_items, m_items, metrics, bs_cache)
            for met in metrics:
                field_vals.setdefault(met, {})[col] = scores.get(
                    met, {"recall": None, "precision": None, "f1": None}
                )

        for met in metrics:
            for campo in COLUMNS:
                v = field_vals[met][campo]
                new_cols[f"{met}_recall_{campo}"].append(v["recall"])
                new_cols[f"{met}_prec_{campo}"].append(v["precision"])
                new_cols[f"{met}_f1_{campo}"].append(v["f1"])

        for met in metrics:
            recalls = [field_vals[met][c]["recall"]    for c in COLUMNS if field_vals[met][c]["recall"]    is not None]
            precs   = [field_vals[met][c]["precision"] for c in COLUMNS if field_vals[met][c]["precision"] is not None]
            f1s     = [field_vals[met][c]["f1"]        for c in COLUMNS if field_vals[met][c]["f1"]        is not None]
            new_cols[f"{met}_recall_TOTAL"].append(round(sum(recalls)/len(recalls), 4) if recalls else None)
            new_cols[f"{met}_prec_TOTAL"].append(round(sum(precs)/len(precs),       4) if precs   else None)
            new_cols[f"{met}_f1_TOTAL"].append(round(sum(f1s)/len(f1s),             4) if f1s     else None)

    df_enriched = df_orig.copy()
    for c in col_names:
        df_enriched[c] = new_cols[c]

    stem = Path(llm_csv).stem
    out_path = out / f"{stem}_con_metricas.csv"
    df_enriched.to_csv(out_path, index=False, encoding="utf-8-sig")
    log.info(f"CSV enriquecido guardado en: {out_path}")

    print("\n" + "=" * 70)
    print("CSV ENRIQUECIDO — columnas añadidas")
    print("=" * 70)
    totals_cols = [c for c in col_names if "TOTAL" in c]
    print(df_enriched[["Id_entrevista"] + totals_cols].to_string(index=False))
    print(f"\n→ {out_path}")
    return df_enriched


# ──────────────────────────────────────────────────────────────────────────────
# CLI
# ──────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    ap = argparse.ArgumentParser(
        description="Métricas NLP (BLEU, chrF++, BERTScore) para el pipeline MEGADAPT",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Ejemplos de uso:

  # Comparar dos corridas + exportar Excel directivo:
  python megadapt_eval_nlp.py \\
      --human "integracion.xlsx" \\
      --llm v5=output_v5/analysis_standardized_v5.csv \\
            v6=output_v6/analysis_standardized_v6.csv \\
      --metrics bleu chrf --xlsx

  # Con BERTScore + Excel (métrica primaria automática = bertscore):
  python megadapt_eval_nlp.py \\
      --human "integracion.xlsx" \\
      --llm v5=output_v5/analysis_standardized_v5.csv \\
      --metrics bleu chrf bertscore \\
      --bertscore-model dccuchile/bert-base-spanish-wwm-cased \\
      --xlsx

  # Enriquecer CSV del pipeline (modo --enrich, sin --xlsx):
  python megadapt_eval_nlp.py \\
      --human "integracion.xlsx" \\
      --llm output_v5/analysis_standardized_v5.csv \\
      --metrics bleu chrf --enrich
""",
    )
    ap.add_argument("--human", required=True,
                    help="Excel de codificación humana (hoja Integracion)")
    ap.add_argument("--llm", nargs="+", required=True,
                    help=(
                        "CSV(s) del pipeline. "
                        "Formato: nombre=ruta.csv  (o solo ruta.csv para nombre automático). "
                        "Se pueden pasar varios para comparar corridas."
                    ))
    ap.add_argument("--out", default="./eval_nlp",
                    help="Directorio de salida (default: ./eval_nlp)")
    ap.add_argument(
        "--metrics", nargs="+", default=["bleu", "chrf"],
        choices=["bleu", "chrf", "bertscore"],
        help="Métricas a calcular (default: bleu chrf).",
    )
    ap.add_argument("--sheet", default="Integracion",
                    help="Nombre de la hoja en el Excel humano (default: Integracion)")
    ap.add_argument(
        "--bertscore-model", default="microsoft/mdeberta-v3-base",
        help=(
            "Modelo HuggingFace para BERTScore:\n"
            "  microsoft/mdeberta-v3-base           (multilingüe, mejor calidad, ~900 MB)\n"
            "  bert-base-multilingual-cased          (ligero, ~700 MB)\n"
            "  dccuchile/bert-base-spanish-wwm-cased (solo español, ~440 MB)"
        ),
    )
    ap.add_argument("--lang", default="es",
                    help="Idioma para BERTScore: 'es', 'en', 'multilingual'")
    ap.add_argument("--device", default="cpu",
                    help="Dispositivo para BERTScore: 'cpu' o 'cuda'")
    ap.add_argument(
        "--enrich", action="store_true",
        help=(
            "Agrega las métricas como columnas nuevas al CSV del pipeline. "
            "Solo funciona con un único --llm. "
            "Genera <nombre>_con_metricas.csv en --out."
        ),
    )
    ap.add_argument(
        "--xlsx", action="store_true",
        help=(
            "Genera un Excel directivo (megadapt_evaluacion.xlsx) con:\n"
            "  Hoja 1 'Benchmarking Modelos': pivot corrida × categoría (F1).\n"
            "  Hoja 2 'Matriz Cruda': datos por entrevista, recall/precision ocultos.\n"
            "  Mapa de calor: Rojo(<0.40) · Amarillo · Verde(>0.85).\n"
            "Solo disponible en modo evaluate (no con --enrich)."
        ),
    )
    ap.add_argument(
        "--primary-metric", default=None,
        choices=["bleu", "chrf", "bertscore"],
        help=(
            "Métrica que se muestra en Hoja 1 del Excel. "
            "Por defecto: bertscore > chrf > bleu (lo que esté disponible)."
        ),
    )
    a = ap.parse_args()

    # Parsear corridas
    runs = {}
    for item in a.llm:
        if "=" in item:
            name, _, path = item.partition("=")
        else:
            name = Path(item).parent.name or "corrida"
            path = item
        runs[name] = path

    if a.enrich:
        if len(runs) > 1:
            ap.error("--enrich solo acepta un único --llm a la vez.")
        if a.xlsx:
            log.warning("--xlsx se ignora en modo --enrich.")
        llm_csv = next(iter(runs.values()))
        enrich_csv(
            a.human, llm_csv, a.out,
            metrics=a.metrics,
            sheet=a.sheet,
            bertscore_model=a.bertscore_model,
            lang=a.lang,
            device=a.device,
        )
    else:
        evaluate(
            a.human, runs, a.out,
            metrics=a.metrics,
            sheet=a.sheet,
            bertscore_model=a.bertscore_model,
            lang=a.lang,
            device=a.device,
            xlsx=a.xlsx,
            primary_metric=a.primary_metric,
        )
